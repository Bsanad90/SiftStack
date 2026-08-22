"""Score the completed live account pull (output/live_account_pull.json) with
the real 4 Pillars of Motivation engine (src/lead_manager.py).

Same engine, same gates, same ranking as score_all_records.py -- this is the
live-data twin, reading the hydrated JSON from src/scripts/live_pull.py
instead of a CSV export. See score_all_records.py for the CSV version and
the notes on why a mapping adapter is needed at all.

Differences from the CSV mapper, both consequences of reading the real API
payload instead of an export:
  * status/lists/tags/phones are native JSON (Title Case status, array lists,
    nested owner.phones), not flattened export columns.
  * there is no `created`/date-added field on the detail record at all (only
    used server-side to bucket the pull), so Timeline scoring here relies
    solely on real event dates (foreclosure_date etc.) with no freshness
    fallback -- a narrower signal than the CSV version's "Created" proxy.

Usage:
    PYTHONPATH=src python src/scripts/score_live_pull.py <input_json> <output_xlsx>
"""
import json
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # src/
import lead_manager  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

DECEASED_LISTS = {"Obituary", "Obituaries Siftmap PRO", "Pre-Probate/Deceased"}
PROBATE_LISTS = {"Probate", "Register of Wills Probates"}
FORECLOSURE_LISTS = {
    "Foreclosure", "Foreclosure Notices", "Pre-Foreclosure",
    "Pre-Foreclosures - Notice of Default", "Siftmap PRO Preforeclosure",
    "Pre-Foreclosures - Lis Pendens",
}
TAX_LISTS = {"Tax Delinquent", "State Tax Liens", "Liens", "Lis Pendens"}
DIVORCE_LISTS = {"Divorce"}

DEAD_STATUSES = {
    "not_interested", "not interested", "dnc", "dead lead", "already sold", "sold", "closed",
    "under_contract", "under contract", "auction date passed", "not related to property",
}

# See score_all_records.py for why this is a precise allowlist rather than a
# substring match, and why blank/unverified structure_type is excluded.
# Caught live on this account: a 229-unit DC co-op ($128M), a 206-unit
# high-rise ($49.7M), a medical clinic ($112.8M) and a children's home
# ($48.7M) all had estimate_value populated and would otherwise have ranked
# top-10 by that field.
SFR_ALLOWED = {"single family residence", "single family residential", "row house", "rural residence"}
MAX_SANE_VALUE = 5_000_000

# See score_all_records.py's identical constant/helper for the full rationale.
# tax_auction_date on this account is a static field, not live-tracked (Ty:
# "it was a static number, not connected to anything"), so an old value
# stops being evidence of imminent urgency but is still worth surfacing --
# just as its own labeled note rather than a false "hot" score.
AUCTION_STALE_FLOOR_DAYS = 30


def auction_signal(raw, today=None):
    """(usable_date_for_the_engine, note_for_display) for a raw auction date string."""
    raw = (raw or "").strip()
    if not raw:
        return "", ""
    try:
        d = datetime.strptime(raw[:10], "%Y-%m-%d")
    except ValueError:
        return raw, ""
    days_until = (d - (today or datetime.now())).days
    if days_until < -AUCTION_STALE_FLOOR_DAYS:
        note = (f"Auction date passed {abs(days_until)} days ago ({raw[:10]}) "
                "-- verify outcome before pursuing")
        return "", note
    return raw, ""


def to_pillar_row(rec):
    lists = set(rec.get("lists") or [])

    deceased = bool(lists & DECEASED_LISTS) or bool(lists & PROBATE_LISTS) \
        or bool(rec.get("last_obituary_date"))

    if lists & FORECLOSURE_LISTS:
        notice_type = "foreclosure"
    elif deceased or (lists & PROBATE_LISTS):
        notice_type = "probate"
    elif lists & TAX_LISTS:
        notice_type = "tax_delinquent"
    elif lists & DIVORCE_LISTS:
        notice_type = "divorce"
    else:
        notice_type = ""

    owner = rec.get("owner") or {}
    company = (owner.get("company") or "").strip()
    owner_name = company or f"{(owner.get('first_name') or '').strip()} {(owner.get('last_name') or '').strip()}".strip()

    addr = rec.get("address") or {}
    vacant = "Y" if addr.get("vacant") else "N"

    auction_usable, auction_note = auction_signal(rec.get("tax_auction_date"))

    mapped = {
        "address": addr.get("street") or "",
        "owner_name": owner_name,
        "notice_type": notice_type,
        "county": addr.get("county") or "",
        "owner_deceased": "yes" if deceased else "no",
        "tax_delinquent_amount": rec.get("tax_delinquent_value") or "",
        "auction_date": auction_usable,
        "date_added": rec.get("foreclosure_date") or "",
        "year_built": rec.get("year") or "",
        "estimated_value": rec.get("estimate_value") or "",
        "sqft": rec.get("sqft") or rec.get("building_sqft") or "",
        "vacant": vacant,
        "equity_percent": rec.get("equity_percent") or "",
        "mls_last_sold_price": rec.get("last_sale_price") or "",
        "decision_maker_name": rec.get("personal_representative") or "",
    }
    return mapped, auction_note


def best_phone(rec):
    owner = rec.get("owner") or {}
    phones = owner.get("phones") or []
    non_wrong = [p for p in phones if (p.get("status") or "").upper() != "WRONG" and p.get("number")]
    pick = (non_wrong or phones or [None])[0]
    if not pick or not pick.get("number"):
        return "", "", ""
    return pick.get("number") or "", pick.get("type") or "", pick.get("status") or ""


def best_email(rec):
    owner = rec.get("owner") or {}
    emails = owner.get("emails") or []
    return emails[0] if emails else ""


def run(in_path, out_path, allowed_types=None, segment_label="Single Family"):
    allowed_types = allowed_types or SFR_ALLOWED
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)

    blob = json.loads(Path(in_path).read_text(encoding="utf-8"))
    entries = blob.get("records") or []

    total = len(entries)
    gate_unhydrated = 0
    gate_status = 0
    gate_no_address = 0
    gate_not_sfr = 0
    gate_insane_value = 0
    gate_no_contact = 0
    quals, extras = [], []

    for entry in entries:
        rec = entry.get("record")
        if not rec:
            gate_unhydrated += 1
            continue

        status = (rec.get("status") or "").strip().lower()
        if status in DEAD_STATUSES:
            gate_status += 1
            continue

        addr = (rec.get("address") or {}).get("street") or ""
        if not addr.strip():
            gate_no_address += 1
            continue

        structure = (rec.get("structure_type") or "").strip().lower()
        if structure not in allowed_types:
            gate_not_sfr += 1
            continue

        try:
            est_v = float(rec.get("estimate_value") or 0)
        except (ValueError, TypeError):
            est_v = 0.0
        if est_v > MAX_SANE_VALUE:
            gate_insane_value += 1
            continue

        phone, phone_type, phone_status = best_phone(rec)
        email = best_email(rec)
        if not phone and not email:
            gate_no_contact += 1
            continue

        mapped, auction_note = to_pillar_row(rec)
        qual = lead_manager.qualify_lead(mapped)
        quals.append(qual)
        a = rec.get("address") or {}
        extras.append({
            "phone": phone, "phone_type": phone_type, "phone_status": phone_status,
            "email": email, "status": rec.get("status") or "",
            "lists": ", ".join(rec.get("lists") or []),
            "mailing_city": a.get("city") or "", "mailing_state": a.get("state") or "",
            "mailing_zip": (a.get("postal_code") or "")[:5],
            "estimated_value": rec.get("estimate_value") or "",
            "equity_percent": rec.get("equity_percent") or "",
            "auction_note": auction_note,
        })

    print(f"Total records in pull:  {total}")
    print(f"Gated - never hydrated: {gate_unhydrated}")
    print(f"Gated - dead status:    {gate_status}")
    print(f"Gated - no address:     {gate_no_address}")
    print(f"Gated - not SFR:        {gate_not_sfr}")
    print(f"Gated - insane value:   {gate_insane_value}")
    print(f"Gated - no contact:     {gate_no_contact}")
    print(f"Scored (in universe):   {len(quals)}")

    hot = sum(1 for q in quals if q.overall_temperature == "hot")
    warm = sum(1 for q in quals if q.overall_temperature == "warm")
    cold = sum(1 for q in quals if q.overall_temperature == "cold")
    print(f"  hot={hot} warm={warm} cold={cold}")

    lead_manager.generate_stabm_report(quals, out_path)

    def est_val(e):
        try:
            return float(str(e["estimated_value"]).replace(",", "").replace("$", "") or 0)
        except ValueError:
            return 0.0

    order = sorted(
        range(len(quals)),
        key=lambda i: (quals[i].hot_count, quals[i].score_total, est_val(extras[i])),
        reverse=True,
    )[:100]

    wb = load_workbook(out_path)
    ws = wb.create_sheet("Top 100 Opportunities", 0)

    hdr_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    hdr_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    title_font = Font(name="Calibri", bold=True, size=16, color="2F5496")
    subtitle_font = Font(name="Calibri", size=11, color="555555")
    thin_border = Border(bottom=Side(style="thin", color="D9D9D9"))
    temp_colors = {
        "hot": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        "warm": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
        "cold": PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
    }

    ws.cell(row=1, column=1, value=f"Top 100 Opportunities -- {segment_label} (Live Account Pull)").font = title_font
    ws.cell(row=2, column=1, value=(
        f"Data pulled {blob.get('started', '?')} -> {blob.get('checkpoint_at', '?')} via apiv2.reisift.io. "
        "Ranked by the 4 Pillars of Motivation engine (src/lead_manager.py): "
        "2+ hot pillars first, then total score (4-12), then estimated value."
    )).font = subtitle_font

    headers = ["Rank", "Owner", "Address", "City", "State", "Zip", "Status", "Est. Value",
               "Equity %", "Hot Pillars", "Score", "Route", "Reason", "Timeline",
               "Condition", "Price", "Phone", "Phone Type", "Phone Status", "Email", "Lists",
               "Auction Note"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill

    for rank, i in enumerate(order, 1):
        q = quals[i]
        e = extras[i]
        row_i = 4 + rank
        vals = [
            rank, q.owner_name, q.address, e["mailing_city"], e["mailing_state"], e["mailing_zip"],
            e["status"], e["estimated_value"], e["equity_percent"], q.hot_count, q.score_total,
            q.route_to, f"{q.reason_pillar.temperature}: {q.reason_pillar.reason}",
            f"{q.timeline_pillar.temperature}: {q.timeline_pillar.reason}",
            f"{q.condition_pillar.temperature}: {q.condition_pillar.reason}",
            f"{q.price_pillar.temperature}: {q.price_pillar.reason}",
            e["phone"], e["phone_type"], e["phone_status"], e["email"], e["lists"],
            e.get("auction_note") or "",
        ]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=row_i, column=col, value=val)
            cell.border = thin_border
        ws.cell(row=row_i, column=10).fill = temp_colors.get(
            "hot" if q.hot_count >= 2 else ("warm" if q.hot_count >= 1 else "cold"), PatternFill()
        )

    widths = [5, 22, 26, 16, 6, 8, 16, 11, 9, 6, 6, 12, 34, 30, 30, 30, 14, 10, 14, 24, 40, 46]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col) if col <= 26 else "A"].width = w
    ws.freeze_panes = "A5"

    ws_a = wb.create_sheet("Audit", 1)
    ws_a.cell(row=1, column=1, value=f"Gate Audit (Live Pull, {segment_label})").font = title_font
    audit_rows = [
        ("Total records in account", total),
        ("Gated: never hydrated (API error/404)", gate_unhydrated),
        ("Gated: dead/converted status", gate_status),
        ("Gated: no property address", gate_no_address),
        (f"Gated: not in segment ({segment_label})", gate_not_sfr),
        (f"Gated: estimated value over ${MAX_SANE_VALUE:,.0f} sanity ceiling", gate_insane_value),
        ("Gated: no phone and no email", gate_no_contact),
        ("Scored (opportunity universe)", len(quals)),
        ("  Hot (2+ hot pillars)", hot),
        ("  Warm", warm),
        ("  Cold", cold),
    ]
    for i, (label, val) in enumerate(audit_rows, 3):
        ws_a.cell(row=i, column=1, value=label)
        ws_a.cell(row=i, column=2, value=val)
    ws_a.column_dimensions["A"].width = 38

    wb.save(out_path)
    print(f"\nSaved: {out_path}")
    return out_path


def main():
    in_path = sys.argv[1] if len(sys.argv) > 1 else "output/live_account_pull.json"
    out_path = sys.argv[2] if len(sys.argv) > 2 else "output/top_100_opportunities_live.xlsx"
    run(in_path, out_path)


if __name__ == "__main__":
    main()
