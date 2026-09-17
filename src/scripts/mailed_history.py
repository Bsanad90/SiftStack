"""Mailed-history index: which properties the vendor mail files say we already mailed.

The DataSift `directmail_attempts` counter is the CRM's only "was mailed" signal and it was
not always advanced -- on the 2026-09-15 obituary cohort, 137 records sat on the Inheritance /
Probate / Register of Wills lists (mailed campaigns) and only 26 of them carried an attempt.
The authoritative record of what went out is the weekly SENT list the mail team files on
Google Drive under `Marketing Lists/<YYYY>/<MM>/<MMDD>/`:

    Probate With PR_<YYYYMMDD>_Mailings.xlsx
    Probate No PR_<YYYYMMDD>_Mailings.xlsx
    Divorce_<YYYYMMDD>_Mailings.xlsx

`vendor_lists_fetch.py --families mailings --years 2026 --since 20260401` downloads them
(read-only on Drive); this script turns them into `output/vendor_lists/mailed_index.json`
keyed by the property address, and screens any mailing CSV against it.

Scope decided by Basem 2026-09-17: 2026 files dated April 1 onward only. The 2025
`Weekly Mailings_*` variants are left out on purpose.

CLOSING THE CYCLE. After a drop is sent, its final CSV must be saved to Drive as
`<Campaign>_<YYYYMMDD>_Mailings.xlsx` under `Marketing Lists/<YYYY>/<MM>/<MMDD>/` (by hand;
the fetcher never writes to Drive). The next fetch + `index` picks it up, and `check`
screens every later list against it.

Usage:
    python src/scripts/mailed_history.py columns            # which header layouts, which columns picked
    python src/scripts/mailed_history.py index [--since 2026-04-01]
    python src/scripts/mailed_history.py check --csv <mailing.csv> [--street-col property_street --zip-col property_zip]

Offline only: nothing here touches DataSift or Drive.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import warnings
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT / "src"), str(ROOT / "src" / "scripts")]

from obituary_dp_batch import sift_key, zip5  # noqa: E402

VENDOR_FILES = ROOT / "output" / "vendor_lists" / "files"
INDEX = ROOT / "output" / "vendor_lists" / "mailed_index.json"
DEFAULT_SINCE = "2026-04-01"

MAILINGS_RE = re.compile(r"_Mailings", re.I)
DATE_IN_NAME = re.compile(r"(\d{8})")
DATE_IN_PATH = re.compile(r"[\\/](\d{4})[\\/](\d{2})[\\/](\d{4})(?:[\\/]|$)")
STREET_LIKE_RE = re.compile(r"^\d+\s+\S+")

# Header names tried IN ORDER for the property (the house the letter is about) and for the
# addressee's own mailing address. Matched case-insensitively after whitespace folding.
PROP_STREET_COLS = ["property address", "property street", "property street address",
                    "site address", "situs address", "address"]
PROP_ZIP_COLS = ["property zip", "property zip code", "property zipcode", "property postal code",
                 "site zip", "situs zip", "zip", "zip code", "zipcode"]
MAIL_STREET_COLS = ["mailing address", "mailing street", "mail address", "mail street",
                    "pr mailing address", "owner address", "owner street", "pr address",
                    "pr address 1"]
MAIL_ZIP_COLS = ["mailing zip", "mailing zip code", "mailing zipcode", "mail zip",
                 "pr mailing zip", "owner zip", "owner zip code", "pr zip", "pr zip code"]


def say(msg: str) -> None:
    print(msg, flush=True)


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def house_key(street: str, zipcode: str) -> str:
    """'<house number> <first street token>|zip5' -- the fuzzy fallback (suffix spelling and
    the unit's spelling -- `Apt 304` vs `#304` -- differ between the vendor file and the CRM).
    Unit IDENTITY is checked separately by `units_compatible`, never ignored."""
    toks = re.sub(r"[^a-z0-9 ]", "", (street or "").lower()).split()
    if len(toks) < 2 or not any(c.isdigit() for c in toks[0]):
        return ""
    return f"{toks[0]} {toks[1]}|{zip5(zipcode)}"


UNIT_RE = re.compile(r"(?:#|\b(?:apt|apartment|unit|ste|suite|bldg|fl|floor|rm|room)\b\.?)\s*#?\s*([a-z0-9-]+)\s*$",
                     re.I)


def unit_of(street: str) -> str:
    """The unit designator at the end of a street line, normalised: 'Apt Ph6' -> 'ph6',
    '# 5-102' -> '5-102', 'Unit 316' -> '316'; '' when the line has none."""
    m = UNIT_RE.search((street or "").strip())
    return m.group(1).lower().strip("-") if m else ""


def units_compatible(a: str, b: str) -> bool:
    """A fuzzy house-key hit is the SAME property only when the unit agrees.

    Measured on the 2026-09-15 obituary cohort: 22 fuzzy hits, 13 of them a DIFFERENT condo
    unit in the same building (`1155 23rd St Nw Apt Ph1L` vs the vendor's `#8b`) -- a
    different property, so not evidence it was mailed. Same unit or no unit on either side
    matches; one side with a unit and the other without is treated as different, so a
    building-level row can never exclude every unit in it.
    """
    ua, ub = unit_of(a), unit_of(b)
    return ua == ub


def file_date(name: str, path: str) -> str:
    """ISO date from the file name's YYYYMMDD, else from the <YYYY>/<MM>/<MMDD> folder."""
    m = DATE_IN_NAME.search(name)
    if m:
        d = m.group(1)
        return f"{d[:4]}-{d[4:6]}-{d[6:]}"
    m = DATE_IN_PATH.search(path)
    if m:
        return f"{m.group(1)}-{m.group(3)[:2]}-{m.group(3)[2:]}"
    return ""


def family_of(name: str) -> str:
    return re.sub(r"[_ ]?\d{8}.*$", "", name).strip() or name


def kind_of(name: str) -> str:
    low = name.lower()
    if "remind" in low:
        return "reminder"
    if "raw" in low:
        return "raw"
    return "sent"


def mailing_files(since_iso: str) -> list[tuple[Path, str]]:
    """[(path, mail_date)] of every downloaded mailing file dated on/after `since_iso`."""
    out = []
    for f in sorted(VENDOR_FILES.rglob("*")):
        if not f.is_file() or f.suffix.lower() not in (".xlsx", ".xlsm", ".xls", ".csv"):
            continue
        if not MAILINGS_RE.search(f.name):
            continue
        d = file_date(f.name, str(f.relative_to(VENDOR_FILES)))
        if not d or d < since_iso:
            continue
        out.append((f, d))
    return out


# ------------------------------------------------------------------ reading a sheet

def _pick(lc: list[str], wanted: list[str]) -> int:
    for w in wanted:
        if w in lc:
            return lc.index(w)
    return -1


def _street_like_col(lc: list[str], rows: list[tuple], exclude: set[int]) -> int:
    """A header containing 'address'/'street' whose first non-empty values look like a
    street line. Used only when none of the named columns exist."""
    for i, c in enumerate(lc):
        if i in exclude or not re.search(r"address|street", c):
            continue
        vals = [_cell(r[i]) for r in rows[:25] if i < len(r) and _cell(r[i])]
        if vals and sum(1 for v in vals if STREET_LIKE_RE.match(v)) >= max(1, len(vals) // 2):
            return i
    return -1


def _zip_near(lc: list[str], street_i: int, exclude: set[int]) -> int:
    """The 'zip' header nearest to (and after) the street column."""
    best, best_d = -1, 10 ** 6
    for i, c in enumerate(lc):
        if i in exclude or "zip" not in c and "postal" not in c:
            continue
        d = (i - street_i) if i > street_i else (street_i - i) + 100
        if d < best_d:
            best, best_d = i, d
    return best


def read_sheet(ws, fname: str) -> tuple[dict, list[dict]]:
    """-> (layout, rows). layout says which columns were picked; rows carry the addresses."""
    it = ws.iter_rows(values_only=True)
    hdr = next(it, None)
    if not hdr:
        return {"headers": [], "prop_street": None, "prop_zip": None,
                "mail_street": None, "mail_zip": None}, []
    cols = [re.sub(r"\s+", " ", _cell(h)) for h in hdr]
    lc = [c.lower() for c in cols]
    body = [r for r in it if any(_cell(v) for v in r)]

    ps = _pick(lc, PROP_STREET_COLS)
    ms = _pick(lc, MAIL_STREET_COLS)
    if ps < 0:
        ps = _street_like_col(lc, body, {ms} if ms >= 0 else set())
    pz = _pick(lc, PROP_ZIP_COLS)
    mz = _pick(lc, MAIL_ZIP_COLS)
    # a generic "zip" must not be paired with the MAILING street when both blocks exist
    if pz >= 0 and mz >= 0 and pz == mz:
        pz = -1
    if ps >= 0 and pz < 0:
        pz = _zip_near(lc, ps, {mz} if mz >= 0 else set())
    if ms >= 0 and mz < 0:
        mz = _zip_near(lc, ms, {pz} if pz >= 0 else set())

    layout = {"headers": cols,
              "prop_street": cols[ps] if ps >= 0 else None,
              "prop_zip": cols[pz] if pz >= 0 else None,
              "mail_street": cols[ms] if ms >= 0 else None,
              "mail_zip": cols[mz] if mz >= 0 else None}
    rows = []
    for r in body:
        def g(i):
            return _cell(r[i]) if 0 <= i < len(r) else ""
        rows.append({"prop_street": g(ps) if ps >= 0 else "",
                     "prop_zip": g(pz) if pz >= 0 else "",
                     "mail_street": g(ms) if ms >= 0 else "",
                     "mail_zip": g(mz) if mz >= 0 else ""})
    return layout, rows


def read_file(f: Path) -> list[tuple[str, dict, list[dict]]]:
    """[(sheet title, layout, rows)] for every sheet in a workbook or the one csv."""
    out = []
    if f.suffix.lower() == ".csv":
        wb = openpyxl.Workbook()
        ws = wb.active
        with open(f, encoding="utf-8-sig", errors="ignore", newline="") as fh:
            for r in csv.reader(fh):
                ws.append(r)
        layout, rows = read_sheet(ws, f.name)
        out.append(("csv", layout, rows))
        return out
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    for ws in wb.worksheets:
        layout, rows = read_sheet(ws, f.name)
        if layout["headers"]:
            out.append((ws.title, layout, rows))
    return out


# ------------------------------------------------------------------ columns

def cmd_columns(a) -> int:
    warnings.filterwarnings("ignore")
    files = mailing_files(a.since)
    if not files:
        say(f"no mailing files dated >= {a.since} under {VENDOR_FILES} -- run "
            "vendor_lists_fetch.py --families mailings first")
        return 2
    layouts: dict[tuple, dict] = {}
    misses, fails = [], []
    for f, d in files:
        rel = str(f.relative_to(VENDOR_FILES))
        try:
            sheets = read_file(f)
        except Exception as exc:  # noqa: BLE001
            fails.append((rel, str(exc)[:120]))
            continue
        for title, layout, rows in sheets:
            k = tuple(layout["headers"])
            L = layouts.setdefault(k, {"layout": layout, "files": [], "rows": 0})
            L["files"].append(f"{rel}[{title}]")
            L["rows"] += len(rows)
            if not layout["prop_street"]:
                misses.append(f"{rel}[{title}]")
    say(f"{len(files)} mailing files dated >= {a.since}, {len(layouts)} distinct header layouts")
    for i, (k, L) in enumerate(sorted(layouts.items(), key=lambda kv: -len(kv[1]["files"])), 1):
        lay = L["layout"]
        say(f"\n[{i}] {len(L['files'])} sheets, {L['rows']} rows")
        say(f"    property: street={lay['prop_street']!r} zip={lay['prop_zip']!r}")
        say(f"    mailing : street={lay['mail_street']!r} zip={lay['mail_zip']!r}")
        say(f"    headers : {' | '.join(k)[:400]}")
        for fn in L["files"][:4]:
            say(f"      e.g. {fn}")
    say(f"\nsheets with NO property street column: {len(misses)}")
    for m in misses:
        say(f"  ! {m}")
    for rel, e in fails:
        say(f"  FAIL {rel}: {e}")
    return 0 if not misses and not fails else 1


# ------------------------------------------------------------------ index

def cmd_index(a) -> int:
    warnings.filterwarnings("ignore")
    files = mailing_files(a.since)
    if not files:
        say(f"no mailing files dated >= {a.since} under {VENDOR_FILES}")
        return 2
    by_prop: dict[str, list[dict]] = defaultdict(list)
    by_house: dict[str, list[dict]] = defaultdict(list)
    by_mail: dict[str, list[dict]] = defaultdict(list)
    fails, no_street_sheets = [], []
    n_rows = n_keyed = 0
    families: Counter = Counter()
    dates = []
    for f, d in files:
        rel = str(f.relative_to(VENDOR_FILES)).replace("\\", "/")
        try:
            sheets = read_file(f)
        except Exception as exc:  # noqa: BLE001
            fails.append({"file": rel, "error": str(exc)[:200]})
            continue
        fam, kind = family_of(f.name), kind_of(f.name)
        families[fam] += 1
        dates.append(d)
        for title, layout, rows in sheets:
            if not layout["prop_street"]:
                no_street_sheets.append(f"{rel}[{title}]")
                continue
            for r in rows:
                n_rows += 1
                hit = {"file": f.name, "drive_path": rel.rsplit("/", 1)[0], "sheet": title,
                       "mail_date": d, "family": fam, "kind": kind,
                       "street": r["prop_street"], "zip": zip5(r["prop_zip"])}
                if r["prop_street"]:
                    n_keyed += 1
                    by_prop[sift_key(r["prop_street"], r["prop_zip"])].append(hit)
                    hk = house_key(r["prop_street"], r["prop_zip"])
                    if hk:
                        by_house[hk].append(hit)
                if r["mail_street"]:
                    by_mail[sift_key(r["mail_street"], r["mail_zip"])].append(
                        dict(hit, mail_street=r["mail_street"], mail_zip=zip5(r["mail_zip"])))
    idx = {"built": datetime.now().isoformat(timespec="seconds"), "since": a.since,
           "files": len(files), "failed_files": fails, "sheets_without_street": no_street_sheets,
           "families": dict(families), "date_range": [min(dates), max(dates)] if dates else [],
           "rows": n_rows, "rows_keyed": n_keyed,
           "by_prop_key": by_prop, "by_house_key": by_house, "by_mail_key": by_mail}
    INDEX.parent.mkdir(parents=True, exist_ok=True)
    INDEX.write_text(json.dumps(idx, indent=0), encoding="utf-8")
    say(f"{len(files)} files ({dict(families)}), dates {idx['date_range']}")
    say(f"rows {n_rows}, with a property street {n_keyed}, distinct properties {len(by_prop)}, "
        f"distinct addressee addresses {len(by_mail)}")
    if no_street_sheets:
        say(f"! {len(no_street_sheets)} sheets had no property street column (see index)")
    for fl in fails:
        say(f"  FAIL {fl['file']}: {fl['error']}")
    say(f"-> {INDEX}")
    return 0 if not fails and not no_street_sheets else 1


# ------------------------------------------------------------------ lookup / check

def load_index(path: Path = INDEX) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def lookup(index: dict, street: str, zipcode: str) -> list[dict]:
    """Hits for a property: exact sift_key first, else house-number+first-token fallback
    (each hit carries `match`: 'exact' | 'fuzzy'). Empty list = never seen in a sent file."""
    if not street:
        return []
    hits = index.get("by_prop_key", {}).get(sift_key(street, zipcode)) or []
    if hits:
        return [dict(h, match="exact") for h in hits]
    hk = house_key(street, zipcode)
    if hk:
        return [dict(h, match="fuzzy") for h in index.get("by_house_key", {}).get(hk) or []
                if units_compatible(street, h.get("street") or "")]
    return []


def describe(hits: list[dict]) -> str:
    """'Probate With PR_20260617_Mailings.xlsx (2026-06-17, sent)' -- oldest first, deduped.
    A fuzzy hit also names the vendor's spelling so the match can be judged in the sheet."""
    seen, parts = set(), []
    for h in sorted(hits, key=lambda h: h["mail_date"]):
        k = (h["file"], h["mail_date"])
        if k in seen:
            continue
        seen.add(k)
        tag = "" if h.get("match", "exact") == "exact" else f", fuzzy: {h.get('street')!r}"
        parts.append(f"{h['file']} ({h['mail_date']}, {h['kind']}{tag})")
    return "; ".join(parts)


def cmd_check(a) -> int:
    if not INDEX.exists():
        say(f"no index at {INDEX} -- run `index` first")
        return 2
    index = load_index()
    src = Path(a.csv)
    rows = list(csv.DictReader(src.open(encoding="utf-8-sig")))
    if not rows:
        say("empty CSV")
        return 2
    for col in (a.street_col, a.zip_col):
        if col not in rows[0]:
            say(f"column {col!r} not in {src.name}; columns: {list(rows[0])}")
            return 2
    out_p = src.with_name(f"{src.stem}_mailed_before.csv")
    n_hit = n_fuzzy = 0
    props_hit: set[str] = set()
    with out_p.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["row", a.street_col, a.zip_col, "match", "mailed_files"])
        for i, r in enumerate(rows, 2):
            hits = lookup(index, r.get(a.street_col) or "", r.get(a.zip_col) or "")
            if not hits:
                continue
            n_hit += 1
            m = hits[0]["match"]
            n_fuzzy += m == "fuzzy"
            props_hit.add(sift_key(r.get(a.street_col) or "", r.get(a.zip_col) or ""))
            w.writerow([i, r.get(a.street_col), r.get(a.zip_col), m, describe(hits)])
    say(f"{len(rows)} rows checked against {index['files']} sent files "
        f"({index['date_range'][0]}..{index['date_range'][1]})")
    say(f"rows previously mailed: {n_hit} ({n_fuzzy} fuzzy) over {len(props_hit)} properties")
    say(f"-> {out_p}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("columns", help="report header layouts and the columns picked per file")
    c.add_argument("--since", default=DEFAULT_SINCE, help="ISO date; files dated earlier are ignored")
    i = sub.add_parser("index", help="build output/vendor_lists/mailed_index.json")
    i.add_argument("--since", default=DEFAULT_SINCE, help="ISO date; files dated earlier are ignored")
    k = sub.add_parser("check", help="screen a mailing CSV against the index")
    k.add_argument("--csv", required=True)
    k.add_argument("--street-col", default="property_street")
    k.add_argument("--zip-col", default="property_zip")
    a = ap.parse_args()
    return {"columns": cmd_columns, "index": cmd_index, "check": cmd_check}[a.cmd](a)


if __name__ == "__main__":
    raise SystemExit(main())
