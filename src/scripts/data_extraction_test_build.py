"""Rebuild the MD/VA tabs of output/data_extraction_test.xlsx from live pulls.

The workbook is Basem's sanity check on the four MD/DC/VA sources (one tab
per source) BEFORE anything is uploaded or skip-traced. It has no builder
until now -- the 2026-08-26/27 version was assembled by hand, which is how
its MD Probates tab ended up with `personal_reps` flattened into one cell and
its MD Legal Notices tab with no publication date. This script:

  * keeps the existing MDDC and VA (Foreclosures) tabs untouched,
  * rebuilds "MD Probates" (Estate Search side) and "MD Legal Notices"
    (Legal Notice Search side) from a `md_register_of_wills_pull.py
    --dump-json` file, with every Personal Representative in its OWN set
    of columns (name / first / last / street / city / state / zip, PR1..PR3),
    the publication date on the Legal Notices tab, and the SDAT-by-street
    property lookup columns (address, current SDAT owner, principal
    residence, confidence, deed date) on both,
  * adds "VA Estate Claims" and "VA Tax Deeds" tabs from the two free
    `va_trustee_sale_pull.py --popular-search 6 / 8` CSVs.

Nothing here touches the account. Usage:

    python src/scripts/data_extraction_test_build.py \\
        --md-json output/md_row_review_08-26-2026.json \\
        --va-estate-claims output/va_estate_claims.csv \\
        --va-tax-deeds output/va_tax_deeds.csv
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from md_register_of_wills_pull import _split_pr_name  # noqa: E402

WORKBOOK = ROOT / "output" / "data_extraction_test.xlsx"

PROPERTY_COLS = [
    "property_street", "property_city", "property_state", "property_zip",
    "property_lookup_confidence", "property_sdat_owner", "property_principal_residence",
    "property_sdat_use", "property_deed_date", "property_deed_book_page",
    "property_lookup_source", "property_lookup_reason",
]


def _pr_columns(n: int) -> list[str]:
    return [f"pr{n}_name", f"pr{n}_first_name", f"pr{n}_last_name",
            f"pr{n}_street", f"pr{n}_city", f"pr{n}_state", f"pr{n}_zip"]


def _pr_values(prs: list[dict], n: int) -> dict:
    pr = prs[n - 1] if len(prs) >= n else {}
    first, last = _split_pr_name(pr.get("name", "")) if pr else ("", "")
    return {
        f"pr{n}_name": pr.get("name", ""), f"pr{n}_first_name": first, f"pr{n}_last_name": last,
        f"pr{n}_street": pr.get("street", ""), f"pr{n}_city": pr.get("city", ""),
        f"pr{n}_state": pr.get("state", ""), f"pr{n}_zip": pr.get("zip", ""),
    }


def _property_values(rec: dict) -> dict:
    out = {c: rec.get(c, "") for c in PROPERTY_COLS}
    if rec.get("property_lookup_confidence") == "NOT_FOUND":
        out["property_street"] = "Not Found (Claude)"
    elif not rec.get("property_lookup_checked"):
        out["property_lookup_confidence"] = "NOT CHECKED"
    return out


PROBATE_COLS = (["county", "estate_number", "estate_type", "status", "filing_date", "date_of_death",
                 "decedent_first_name", "decedent_last_name", "source", "pr_count"]
                + _pr_columns(1) + _pr_columns(2) + _pr_columns(3)
                + ["attorney_name", "attorney_street", "attorney_city", "attorney_state", "attorney_zip"]
                + PROPERTY_COLS)

NOTICE_COLS = (["published_on", "county", "estate_number", "date_of_death",
                "decedent_first_name", "decedent_last_name", "source", "pr_count"]
               + _pr_columns(1) + _pr_columns(2) + _pr_columns(3)
               + PROPERTY_COLS)


def probate_rows(records: list[dict]) -> list[dict]:
    rows = []
    for rec in records:
        if rec.get("source") not in ("estate_search", "both"):
            continue
        prs = rec.get("personal_reps") or []
        att = rec.get("attorney") or {}
        row = {
            "county": rec.get("county", ""), "estate_number": rec.get("estate_number", ""),
            "estate_type": rec.get("estate_type", ""), "status": rec.get("status", ""),
            "filing_date": rec.get("filing_date") or rec.get("date_of_filing", ""),
            "date_of_death": rec.get("date_of_death", ""),
            "decedent_first_name": rec.get("decedent_first_name", ""),
            "decedent_last_name": rec.get("decedent_last_name", ""),
            "source": rec.get("source", ""), "pr_count": len(prs),
            "attorney_name": att.get("name", ""), "attorney_street": att.get("street", ""),
            "attorney_city": att.get("city", ""), "attorney_state": att.get("state", ""),
            "attorney_zip": att.get("zip", ""),
        }
        for n in (1, 2, 3):
            row.update(_pr_values(prs, n))
        row.update(_property_values(rec))
        rows.append(row)
    return rows


def notice_rows(records: list[dict]) -> list[dict]:
    rows = []
    for rec in records:
        if rec.get("source") not in ("legal_notice", "both"):
            continue
        prs = rec.get("personal_reps") or []
        row = {
            "published_on": rec.get("legal_notice_published_on") or rec.get("published_on", ""),
            "county": rec.get("county", ""), "estate_number": rec.get("estate_number", ""),
            "date_of_death": rec.get("date_of_death", ""),
            "decedent_first_name": rec.get("decedent_first_name", ""),
            "decedent_last_name": rec.get("decedent_last_name", ""),
            "source": rec.get("source", ""), "pr_count": len(prs),
        }
        for n in (1, 2, 3):
            row.update(_pr_values(prs, n))
        row.update(_property_values(rec))
        rows.append(row)
    return rows


def csv_rows(path: Path, drop: tuple[str, ...] = ("full_text",)) -> tuple[list[str], list[dict]]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        cols = [c for c in (reader.fieldnames or []) if c not in drop]
        rows = [{c: r.get(c, "") for c in cols} for r in reader]
    return cols, rows


HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def write_sheet(wb: openpyxl.Workbook, title: str, cols: list[str], rows: list[dict], note: str = "") -> None:
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for r in rows:
        ws.append([_cell(r.get(c, "")) for c in cols])
    if not rows:
        ws.append([f"(no rows) {note}".strip()])
    ws.freeze_panes = "A2"
    for i, col in enumerate(cols, start=1):
        longest = max([len(str(col))] + [len(str(r.get(col, "") or "")) for r in rows])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 60)
    if note:
        ws.cell(row=1, column=len(cols) + 2, value=note).font = Font(italic=True, color="7F7F7F")


def _cell(v):
    if isinstance(v, (list, dict)):
        return json.dumps(v)
    if isinstance(v, str) and len(v) > 32000:
        return v[:32000]
    return v


def save(wb: openpyxl.Workbook, path: Path) -> Path:
    try:
        wb.save(path)
        return path
    except PermissionError:
        pending = path.with_name(f"_PENDING_{path.name}")
        wb.save(pending)
        print(f"  WARNING: {path.name} is open in Excel -- wrote {pending.name}; close the file and swap it in.")
        return pending


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook", default=str(WORKBOOK))
    ap.add_argument("--md-json", default="", help="md_register_of_wills_pull.py --dump-json output")
    ap.add_argument("--va-estate-claims", default="", help="va_trustee_sale_pull.py --popular-search 6 CSV")
    ap.add_argument("--va-tax-deeds", default="", help="va_trustee_sale_pull.py --popular-search 8 CSV")
    ap.add_argument("--limit", type=int, default=0, help="cap rows per rebuilt tab (0 = all)")
    args = ap.parse_args()

    path = Path(args.workbook)
    wb = openpyxl.load_workbook(path) if path.exists() else openpyxl.Workbook()
    if not path.exists():
        del wb[wb.sheetnames[0]]

    def cap(rows):
        return rows[: args.limit] if args.limit else rows

    if args.md_json:
        dump = json.loads(Path(args.md_json).read_text(encoding="utf-8"))
        run = dump.get("run", {})
        note = (f"pull {run.get('date_from')}..{run.get('date_to')} counties={','.join(run.get('counties', []))} "
                f"site data as of {run.get('site_cutoff')}; committed={run.get('committed')}")
        p_rows = cap(probate_rows(dump["records"]))
        n_rows = cap(notice_rows(dump["records"]))
        write_sheet(wb, "MD Probates", PROBATE_COLS, p_rows, note)
        write_sheet(wb, "MD Legal Notices", NOTICE_COLS, n_rows, note)
        print(f"MD Probates: {len(p_rows)} row(s); MD Legal Notices: {len(n_rows)} row(s)")
        found = sum(1 for r in p_rows + n_rows if r["property_lookup_confidence"] in ("HIGH", "MEDIUM", "LOW"))
        print(f"  property address resolved on {found} of {len(p_rows) + len(n_rows)} tab rows")

    for flag, title in ((args.va_estate_claims, "VA Estate Claims"), (args.va_tax_deeds, "VA Tax Deeds")):
        if not flag:
            continue
        cols, rows = csv_rows(Path(flag))
        rows = cap(rows)
        write_sheet(wb, title, cols, rows, note=f"from {Path(flag).name}")
        print(f"{title}: {len(rows)} row(s)")

    out = save(wb, path)
    print(f"Wrote {out}  tabs: {wb.sheetnames}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
