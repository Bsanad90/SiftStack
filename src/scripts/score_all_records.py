"""Score the full account export with the real 4 Pillars of Motivation engine
(src/lead_manager.py) and produce a Top 100 Opportunities workbook.

This does NOT reimplement scoring. It maps the REISift "All Records" export
(different column names/shape than the outbound upload CSV lead_manager.py was
built against) into the exact row schema qualify_lead() expects, then calls
the unmodified engine.

Usage:
    PYTHONPATH=src python src/scripts/score_all_records.py <input_csv> <output_xlsx>
"""
import csv
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # src/
import lead_manager  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

# ── Lists -> notice_type / deceased inference (REISift export has no
# notice_type/owner_deceased columns; TN pipeline's vocabulary is inferred
# from list membership instead) ─────────────────────────────────────────
DECEASED_LISTS = {"Obituary", "Obituaries Siftmap PRO", "Pre-Probate/Deceased"}
PROBATE_LISTS = {"Probate", "Register of Wills Probates"}
FORECLOSURE_LISTS = {
    "Foreclosure", "Foreclosure Notices", "Pre-Foreclosure",
    "Pre-Foreclosures - Notice of Default", "Siftmap PRO Preforeclosure",
    "Pre-Foreclosures - Lis Pendens",
}
TAX_LISTS = {"Tax Delinquent", "State Tax Liens", "Liens", "Lis Pendens"}
DIVORCE_LISTS = {"Divorce"}

# Records in these statuses are already dead/converted/resolved -- scoring
# them as a fresh "opportunity" would be wrong. This gate sits in front of
# the 4 Pillars engine; it is not part of the engine itself.
DEAD_STATUSES = {
    "not_interested", "dnc", "dead lead", "already sold", "sold", "closed",
    "under_contract", "auction date passed", "not related to property",
}

# Precise allowlist, not a substring match. Deliberately narrower than a
# "contains residential" check: this account's real structure_type values
# include "Cooperative Building (Residential)", "Commercial Office/
# Residential (Mixed Use)" and "Residential Parking Garage", all of which
# contain "residential" as a substring. A loose match let a 229-unit DC co-op
# ($128M), a 206-unit high-rise ($49.7M), a medical clinic ($112.8M) and a
# children's home ($48.7M) sail into the first live run's top 10, because
# the estimate_value field carries the whole building's value, not a unit's.
# Mirrors the "single family only" buy box already standing for the Knox
# pipeline (see CLAUDE.md). Blank/unverified structure_type is excluded
# rather than assumed safe -- same rule the rest of this codebase applies to
# missing value/equity data.
SFR_ALLOWED = {"single family residence", "single family residential", "row house", "rural residence"}
MAX_SANE_VALUE = 5_000_000  # backstop even for an allowed type; generous for DC-metro luxury

# How far past a scheduled auction date the engine's own hot/warm math is
# still trusted (Ty: these dates are static, not live-tracked, so an old one
# is meaningless -- but a recent one may just be an unconfirmed postponement,
# still worth calling). Beyond this floor the date is withheld from the
# engine (falls through to the freshness/filing-date branch instead) and
# surfaced as its own "Auction Note" column instead of silently vanishing
# into an indistinguishable "cold: no urgency" bucket.
AUCTION_STALE_FLOOR_DAYS = 30


def split_lists(raw):
    return set(x.strip() for x in (raw or "").split(",") if x.strip())


def auction_signal(raw, today=None):
    """(usable_date_for_the_engine, note_for_display) for a raw auction date string."""
    raw = (raw or "").strip()
    if not raw:
        return "", ""
    try:
        d = datetime.strptime(raw[:10], "%Y-%m-%d")
    except ValueError:
        return raw, ""  # unrecognized format: pass through untouched, no note
    days_until = (d - (today or datetime.now())).days
    if days_until < -AUCTION_STALE_FLOOR_DAYS:
        note = (f"Auction date passed {abs(days_until)} days ago ({raw[:10]}) "
                "-- verify outcome before pursuing")
        return "", note
    return raw, ""


def to_pillar_row(rec):
    """Map one REISift export row into the dict qualify_lead() expects."""
    lists = split_lists(rec.get("Lists"))

    deceased = bool(lists & DECEASED_LISTS) or bool(lists & PROBATE_LISTS) \
        or bool((rec.get("Last obituary date") or "").strip())

    if lists & FORECLOSURE_LISTS:
        notice_type = "foreclosure"
    elif deceased or (lists & PROBATE_LISTS):
        notice_type = "probate"
    elif lists & TAX_LISTS:
        notice_type = "tax_delinquent"
    elif lists & DIVORCE_LISTS:
        notice_type = "divorce"
    else:
        notice_type = ""

    business = (rec.get("Business Name") or "").strip()
    owner_name = business or f"{(rec.get('First Name') or '').strip()} {(rec.get('Last Name') or '').strip()}".strip()

    vacant = "Y" if (rec.get("Property vacant") or "").strip().lower() == "true" else "N"

    # "Tax auction date" is a genuine scheduled sale date -> feeds the
    # engine's forward-looking "auction in N days" branch. "Foreclosure
    # date" in this export is the notice/filing date (often long past), NOT
    # a future auction date -- feeding IT into the same slot made a
    # 709-day-old filing score as "hot: auction in -709 days" because the
    # engine's days_until <= 30 check has no floor. A genuinely stale
    # scheduled date has the identical problem (see auction_signal above),
    # so both are withheld past the floor rather than handed to the engine.
    auction_usable, auction_note = auction_signal(rec.get("Tax auction date"))

    mapped = {
        "address": rec.get("Property address") or "",
        "owner_name": owner_name,
        "notice_type": notice_type,
        "county": rec.get("Property county") or "",
        "owner_deceased": "yes" if deceased else "no",
        "tax_delinquent_amount": rec.get("Tax delinquent value") or "",
        "auction_date": auction_usable,
        "date_added": rec.get("Foreclosure date") or rec.get("Created") or "",
        "year_built": rec.get("Year") or "",
        "estimated_value": rec.get("Estimated value") or "",
        "sqft": rec.get("Sqft") or "",
        "vacant": vacant,
        "equity_percent": rec.get("Equity percent") or "",
        "mls_last_sold_price": rec.get("Last sale price") or "",
        "decision_maker_name": rec.get("Personal representative") or "",
    }
    return mapped, auction_note


def best_phone(rec):
    for i in range(1, 6):
        p = (rec.get(f"Phone {i}") or "").strip()
        if p:
            return p, (rec.get(f"Phone Type {i}") or "").strip(), (rec.get(f"Phone Status {i}") or "").strip()
    return "", "", ""


def main():
    in_path = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\pc\Desktop\Galal Development\All Records 8.21.2026.csv"
    out_path = sys.argv[2] if len(sys.argv) > 2 else r"C:\Users\pc\Desktop\Galal Development\Sift Stack\output\top_100_opportunities.xlsx"
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)

    total = 0
    gate_status = 0
    gate_no_contact = 0
    gate_no_address = 0
    gate_not_sfr = 0
    gate_insane_value = 0
    quals = []
    extras = []  # parallel list: dict of display extras per surviving row

    with open(in_path, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for rec in reader:
            total += 1

            status = (rec.get("Status") or "").strip().lower()
            if status in DEAD_STATUSES:
                gate_status += 1
                continue

            addr = (rec.get("Property address") or "").strip()
            if not addr:
                gate_no_address += 1
                continue

            structure = (rec.get("Structure type") or "").strip().lower()
            if structure not in SFR_ALLOWED:
                gate_not_sfr += 1
                continue

            try:
                est_v = float(str(rec.get("Estimated value") or "0").replace(",", "").replace("$", ""))
            except ValueError:
                est_v = 0.0
            if est_v > MAX_SANE_VALUE:
                gate_insane_value += 1
                continue

            phone, phone_type, phone_status = best_phone(rec)
            email = (rec.get("Email 1") or "").strip()
            if not phone and not email:
                gate_no_contact += 1
                continue

            mapped, auction_note = to_pillar_row(rec)
            qual = lead_manager.qualify_lead(mapped)
            quals.append(qual)
            extras.append({
                "phone": phone, "phone_type": phone_type, "phone_status": phone_status,
                "email": email, "status": (rec.get("Status") or "").strip(),
                "lists": (rec.get("Lists") or "").strip(),
                "mailing_city": rec.get("Property city") or "",
                "mailing_state": rec.get("Property state") or "",
                "mailing_zip": rec.get("Property zip5") or rec.get("Property zip") or "",
                "estimated_value": rec.get("Estimated value") or "",
                "equity_percent": rec.get("Equity percent") or "",
                "auction_note": auction_note,
            })

    print(f"Total records:          {total}")
    print(f"Gated - dead status:    {gate_status}")
    print(f"Gated - no address:     {gate_no_address}")
    print(f"Gated - not SFR:        {gate_not_sfr}")
    print(f"Gated - insane value:   {gate_insane_value}")
    print(f"Gated - no contact:     {gate_no_contact}")
    print(f"Scored (in universe):   {len(quals)}")

    hot = sum(1 for q in quals if q.overall_temperature == "hot")
    warm = sum(1 for q in quals if q.overall_temperature == "warm")
    cold = sum(1 for q in quals if q.overall_temperature == "cold")
    print(f"  hot={hot} warm={warm} cold={cold}")

    # Standard STABM deliverable, generated by the unmodified engine.
    lead_manager.generate_stabm_report(quals, out_path)

    # Rank for Top 100: 2+ hot pillars first (the engine's own "closer queue"
    # bar), then total score, then estimated value as a business tiebreaker.
    def est_val(e):
        try:
            return float(str(e["estimated_value"]).replace(",", "").replace("$", "") or 0)
        except ValueError:
            return 0.0

    order = sorted(
        range(len(quals)),
        key=lambda i: (quals[i].hot_count, quals[i].score_total, est_val(extras[i])),
        reverse=True,
    )[:100]

    # Re-open the saved workbook and insert a Top 100 sheet at the front.
    wb = load_workbook(out_path)
    ws = wb.create_sheet("Top 100 Opportunities", 0)

    hdr_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    hdr_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    title_font = Font(name="Calibri", bold=True, size=16, color="2F5496")
    subtitle_font = Font(name="Calibri", size=11, color="555555")
    thin_border = Border(bottom=Side(style="thin", color="D9D9D9"))
    temp_colors = {
        "hot": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        "warm": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
        "cold": PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),
    }

    ws.cell(row=1, column=1, value="Top 100 Opportunities").font = title_font
    ws.cell(row=2, column=1, value=(
        "Ranked by the 4 Pillars of Motivation engine (src/lead_manager.py): "
        "2+ hot pillars first, then total score (4-12), then estimated value."
    )).font = subtitle_font

    headers = ["Rank", "Owner", "Address", "City", "State", "Zip", "Status", "Est. Value",
               "Equity %", "Hot Pillars", "Score", "Route", "Reason", "Timeline",
               "Condition", "Price", "Phone", "Phone Type", "Phone Status", "Email", "Lists",
               "Auction Note"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill

    for rank, i in enumerate(order, 1):
        q = quals[i]
        e = extras[i]
        row_i = 4 + rank
        vals = [
            rank, q.owner_name, q.address, e["mailing_city"], e["mailing_state"], e["mailing_zip"],
            e["status"], e["estimated_value"], e["equity_percent"], q.hot_count, q.score_total,
            q.route_to, f"{q.reason_pillar.temperature}: {q.reason_pillar.reason}",
            f"{q.timeline_pillar.temperature}: {q.timeline_pillar.reason}",
            f"{q.condition_pillar.temperature}: {q.condition_pillar.reason}",
            f"{q.price_pillar.temperature}: {q.price_pillar.reason}",
            e["phone"], e["phone_type"], e["phone_status"], e["email"], e["lists"],
            e.get("auction_note") or "",
        ]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=row_i, column=col, value=val)
            cell.border = thin_border
        ws.cell(row=row_i, column=10).fill = temp_colors.get(
            "hot" if q.hot_count >= 2 else ("warm" if q.hot_count >= 1 else "cold"), PatternFill()
        )

    widths = [5, 22, 26, 16, 6, 8, 16, 11, 9, 6, 6, 12, 34, 30, 30, 30, 14, 10, 14, 24, 40, 46]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col) if col <= 26 else "A"].width = w
    ws.freeze_panes = "A5"

    # Audit sheet
    ws_a = wb.create_sheet("Audit", 1)
    ws_a.cell(row=1, column=1, value="Gate Audit").font = title_font
    audit_rows = [
        ("Total records in export", total),
        ("Gated: dead/converted status", gate_status),
        ("Gated: no property address", gate_no_address),
        ("Gated: not single family (structure_type allowlist)", gate_not_sfr),
        (f"Gated: estimated value over ${MAX_SANE_VALUE:,.0f} sanity ceiling", gate_insane_value),
        ("Gated: no phone and no email", gate_no_contact),
        ("Scored (opportunity universe)", len(quals)),
        ("  Hot (2+ hot pillars)", hot),
        ("  Warm", warm),
        ("  Cold", cold),
    ]
    for i, (label, val) in enumerate(audit_rows, 3):
        ws_a.cell(row=i, column=1, value=label)
        ws_a.cell(row=i, column=2, value=val)
    ws_a.column_dimensions["A"].width = 34

    wb.save(out_path)
    print(f"\nSaved: {out_path}")


if __name__ == "__main__":
    main()
