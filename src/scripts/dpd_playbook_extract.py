"""Extract one jurisdiction's doors-per-deal playbook data into a machine-readable file.

    python src/scripts/dpd_playbook_extract.py --fips 11001 \
        --xlsx "output/market_research/11001-District-of-Columbia-DC-doors-per-deal.xlsx" --fetch

Two sources, same data, either or both:

  * The SINGLE-COUNTY workbook exported from learn.datasift.ai/county-list-playbook#<fips>
    with the page's "Download Excel" button. Since the 2026-09 page release the export is
    the WORK ORDER layout: sheets Overview / Work Order / Obituary / First to Market /
    Top ZIPs / Price Bands / Read Me. The old layout (Priority Ranking / AI Score /
    Benchmarks) is gone from the site and is not accepted: a file without a Work Order
    sheet is refused rather than half-parsed.
  * The playbook page's own JSON shards, free and unauthenticated, keyed by the 2-digit
    state FIPS: /county-data/<st>.json, /ftm-data/<st>.json, /capture-data/<st>.json.
    `--fetch` downloads them. capture-data is what the Work Order is drawn from, and it
    holds TWO ladders per county: `ai` (the default on screen, and what the export
    carries) and `lists` (the same pull order with AI rungs removed, so the deals AI
    would have claimed are re-credited to the lists below it). With a workbook present
    the shards are a CROSS-CHECK and add what the sheet lacks: each rung's standalone
    `list_size`, its stable signal `keys` (is_absentee_owners ...) and the no-AI ladder
    as `rows_no_ai`. Without a workbook they are the source.

READ THE NUMBERS AS A PULL ORDER. A Work Order rung's doors/deal and lift are MARGINAL:
they count only the deals the rungs above it do not already reach, so they price the
next pull, not the list on its own. `list_size` (from the shard) is the standalone list.

Writes data/dpd_playbook_<fips>.json. Read-only against everything.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import urllib.request
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))

from dpd.jurisdictions import BY_FIPS  # noqa: E402

PLAYBOOK = "https://learn.datasift.ai"
SHARDS = ("county-data", "ftm-data", "capture-data")

LADDER_HEADER = ["Rung", "List to pull", "Type", "Plan", "Doors per deal (this rung)",
                 "Vs baseline", "New deals reached", "Typical gross", "Market share after %",
                 "Total live doors", "Caveats"]
ZIP_HEADER = ["Rank", "ZIP", "Rating (1-5)", "Deals (6mo)", "Typical gross", "Doors/deal",
              "Momentum (x)"]
PRICE_HEADER = ["Value Band", "Doors/Deal", "Deals (6mo)", "Live SFR Supply", "Note"]
FTM_HEADER = ["Priority", "Data Type", "Office / Official", "Phone", "Source URL",
              "Office Address", "Records / FOIA", "Access", "Updates", "Verified?", "Notes"]
FTM_KEYS = ["priority", "data_type", "office_official", "phone", "source_url",
            "office_address", "records_foia", "access", "updates", "verified", "notes"]
REQUIRED_SHEETS = {"Overview", "Work Order", "First to Market", "Top ZIPs"}


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, str) and v.strip().lower() in ("n/a", "-", ""):
        return None
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def _rows(ws) -> list[tuple]:
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def _c(r: tuple, i: int):
    """Cell i, or None past the end of the row.

    openpyxl read_only mode TRIMS trailing empty cells, so a row whose last columns are
    blank comes back short (a zero-deal band with no gross, a rung with no caveat). Every
    fixed-index read of a data row goes through here.
    """
    return r[i] if i < len(r) else None


def _after_header(rows: list[tuple], header: list[str]) -> list[tuple]:
    for i, r in enumerate(rows):
        cells = [str(c).strip() if c is not None else "" for c in r[:len(header)]]
        if cells == header:
            return [x for x in rows[i + 1:] if any(c not in (None, "") for c in x)]
    raise SystemExit(f"header not found: {header[:3]}... (has the export layout changed again?)")


def _signals(label: str) -> list[str]:
    return [s.strip() for s in (label or "").split(" + ")]


def _rank_within_priority(rows: list[dict]) -> None:
    counters: dict = {}
    for row in rows:
        p = row["priority"]
        counters[p] = counters.get(p, 0) + 1
        row["rank"] = counters[p]


def parse_ladder(rows_in: list[tuple]) -> tuple[dict, list[dict], list[str]]:
    """The Work Order sheet -> (baseline, rungs, footnotes).

    Layout: header, a rung-0 baseline row, then 'Priority N' section rows each followed by
    numbered rungs, then free-text footnotes. Priority comes from the section a rung sits
    in; the sheet has no priority column.
    """
    baseline, rows, notes, pri, prev_doors = {}, [], [], None, 0
    for r in _after_header(rows_in, LADDER_HEADER):
        r = tuple(r) + (None,) * (len(LADDER_HEADER) - len(r))
        head = _c(r, 0)
        if isinstance(head, str):
            m = re.match(r"^Priority\s+(\d)$", head.strip())
            if m:
                pri = int(m.group(1))
            else:
                notes.append(head.strip())
            continue
        rung = _num(head)
        if rung is None:
            continue
        if int(rung) == 0:
            baseline = {"dpd": _num(_c(r, 4)), "deals": _num(_c(r, 6)),
                        "typical_gross": _num(_c(r, 7)), "list_size": _num(_c(r, 9))}
            continue
        total = _num(_c(r, 9))
        label = str(_c(r, 1)).strip()
        rtype = _c(r, 2)
        rows.append({
            "rung": int(rung), "priority": pri, "signal": label, "signals": _signals(label),
            "type": rtype, "plan": _c(r, 3), "ai": rtype == "AI score",
            "dpd": _num(_c(r, 4)), "lift": _num(_c(r, 5)), "deals": _num(_c(r, 6)),
            "typical_gross": _num(_c(r, 7)), "market_share_after_pct": _num(_c(r, 8)),
            "total_live_doors": total,
            "doors_added": (total - prev_doors) if total is not None else None,
            "caveats": (str(_c(r, 10)).strip() if _c(r, 10) else None),
            # standalone list size and signal keys are not on the sheet; --fetch fills them
            "list_size": None, "keys": None,
        })
        if total is not None:
            prev_doors = total
    _rank_within_priority(rows)
    return baseline, rows, notes


def parse_obituary(rows_in: list[tuple]) -> dict:
    out: dict = {"narrative": [], "timeline": []}
    labels = {"Flagged homes here": "flagged_homes", "New filings a month": "new_per_month",
              "Flagged nationwide": "flagged_nationwide", "Your head start": "head_start",
              "Modeled doors per deal, optimistic": "modeled_dpd_optimistic",
              "Modeled doors per deal, base case": "modeled_dpd_base",
              "Modeled doors per deal, conservative": "modeled_dpd_conservative"}
    in_timeline = False
    for r in rows_in:
        k = _c(r, 0)
        if k in (None, ""):
            continue
        k = str(k).strip()
        if k == "Stage":
            in_timeline = True
            continue
        if in_timeline:
            out["timeline"].append({"stage": k, "what_happens": str(_c(r, 1) or "").strip()})
        elif k in labels:
            v = _c(r, 1)
            out[labels[k]] = _num(v) if _num(v) is not None else v
        elif _c(r, 1) in (None, "") and len(k) > 60:
            out["narrative"].append(k)
    return out


def parse_xlsx(path: Path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    missing = REQUIRED_SHEETS - set(wb.sheetnames)
    if missing:
        wb.close()
        raise SystemExit(
            f"{path.name} lacks sheets {sorted(missing)}. This parser reads only the Work Order "
            "layout the playbook page exports since 2026-09 (Overview / Work Order / Obituary / "
            "First to Market / Top ZIPs / Price Bands / Read Me). Re-export from "
            "learn.datasift.ai/county-list-playbook with the Download Excel button.")

    out: dict = {"source_xlsx": path.name, "format": "work_order"}

    # Overview: title row + 2-column facts. Labels are kept verbatim (the page's own
    # wording, e.g. "Baseline doors/deal", "SFR supply") and consumers read them as such.
    ov = _rows(wb["Overview"])
    out["title"] = str(ov[0][0]) if ov and ov[0] else None
    m = re.search(r"FIPS\s+(\d{5})", out["title"] or "")
    out["fips_in_title"] = m.group(1) if m else None
    out["overview"] = {str(_c(r, 0)).strip(): _c(r, 1) for r in ov
                       if len(r) >= 2 and _c(r, 0) not in (None, "") and _c(r, 1) not in (None, "")}

    baseline, rows, notes = parse_ladder(_rows(wb["Work Order"]))
    out["baseline"] = baseline
    out["rows"] = rows
    out["ladder_notes"] = notes
    mm = re.search(r"measured (\d{4}-\d{2}-\d{2})", " ".join(notes))
    out["ladder_measured"] = mm.group(1) if mm else None

    if "Obituary" in wb.sheetnames:
        out["obituary"] = parse_obituary(_rows(wb["Obituary"]))

    zips = []
    for r in _after_header(_rows(wb["Top ZIPs"]), ZIP_HEADER):
        zips.append({"rank": _num(_c(r, 0)), "zip": str(_c(r, 1)).strip(),
                     "stars": int(_num(_c(r, 2)) or 0), "deals": _num(_c(r, 3)),
                     "typical_gross": _num(_c(r, 4)), "dpd": _num(_c(r, 5)),
                     "momentum": _num(_c(r, 6))})
    out["top_zips"] = zips

    if "Price Bands" in wb.sheetnames:
        out["price_bands"] = [
            {"band": _c(r, 0), "dpd": _num(_c(r, 1)), "deals": _num(_c(r, 2)),
             "live_sfr_supply": _num(_c(r, 3)), "note": _c(r, 4)}
            for r in _after_header(_rows(wb["Price Bands"]), PRICE_HEADER)
            if _num(_c(r, 2)) is not None]

    # First to Market: 11 columns, then narrative rows (gap paragraph, how-to-read).
    ftm_rows, narrative = [], []
    for r in _after_header(_rows(wb["First to Market"]), FTM_HEADER):
        r = tuple(r) + (None,) * (len(FTM_HEADER) - len(r))
        if _c(r, 1) in (None, "") and _c(r, 0):
            narrative.append(str(_c(r, 0)).strip())
            continue
        ftm_rows.append({k: (str(v).strip() if v is not None else None)
                         for k, v in zip(FTM_KEYS, r)})
    out["ftm"] = ftm_rows
    out["ftm_narrative"] = narrative
    gap = next((n for n in narrative if n.startswith("Data gaps")), "")
    low = re.search(r"for:\s*(.+?)\.\s", gap + " ")
    absent = re.search(r"Also not in SiftMap here:\s*(.+?)\.?$", gap)
    out["siftmap_gaps"] = {
        "low_or_zero": [s.strip() for s in low.group(1).split(",")] if low else [],
        "absent": [s.strip() for s in absent.group(1).split(",")] if absent else [],
        "verbatim": gap or None,
    }

    if "Read Me" in wb.sheetnames:
        out["read_me"] = {str(_c(r, 0)).strip(): str(_c(r, 1)).strip()
                          for r in _rows(wb["Read Me"]) if len(r) >= 2 and _c(r, 0) and _c(r, 1)}
    wb.close()   # read_only keeps the file handle open until closed (Windows then locks it)
    return out


def fetch_shards(state_fips2: str, dest: Path) -> dict:
    got = {}
    for k in SHARDS:
        url = f"{PLAYBOOK}/{k}/{state_fips2}.json"
        with urllib.request.urlopen(url, timeout=60) as resp:  # noqa: S310
            got[k] = json.loads(resp.read().decode("utf-8"))
    dest.write_text(json.dumps(got, indent=1), encoding="utf-8")
    return got


def rows_from_capture(variant: list[dict], base_dpd: float | None) -> list[dict]:
    """One capture-data ladder variant -> the same row shape parse_ladder() emits."""
    rows = []
    for i, r in enumerate(variant or []):
        label = r.get("label") or ""
        marg = r.get("marg_dpd")
        rows.append({
            "rung": i + 1, "priority": r.get("priority"), "signal": label,
            "signals": _signals(label),
            "type": "AI score" if r.get("ai") else ("Stack" if " + " in label else "List"),
            "plan": r.get("plan"), "ai": bool(r.get("ai")),
            "dpd": marg,
            "lift": (round(base_dpd / marg, 1) if base_dpd and marg else None),
            "deals": r.get("marg_deals"), "typical_gross": r.get("median_gross"),
            "market_share_after_pct": r.get("cum_share"),
            "total_live_doors": r.get("cum_supply"), "doors_added": r.get("marg_supply"),
            "caveats": None, "list_size": r.get("list_size"), "keys": r.get("keys"),
            "est": r.get("est"), "thin": r.get("thin"),
        })
    _rank_within_priority(rows)
    return rows


def _shard_extra(c: dict, cap: dict) -> dict:
    return {"coverage": c.get("coverage"), "obit_supply": c.get("obit_supply"),
            "ai_demand": c.get("ai_demand"),
            "ftm_verdicts": c.get("ftm"), "zips_detail": c.get("zips"),
            "trend": c.get("trend"), "capture_ladder": cap}


def _capture_baseline(cap: dict) -> dict:
    deals, supply = cap.get("deals"), cap.get("base_supply")
    return {"deals": deals, "list_size": supply,
            "dpd": (round(supply / deals, 1) if deals and supply else None),
            "typical_gross": cap.get("baseline_gross")}


def merge_shard(out: dict, shards: dict, fips: str) -> None:
    c = (shards.get("county-data") or {}).get(fips) or {}
    cap = (shards.get("capture-data") or {}).get(fips) or {}
    if not cap:
        out["shard_check"] = {"status": "county not in capture-data shard"}
        return
    check: dict = {"status": "ok", "mismatches": []}
    cb = _capture_baseline(cap)
    ov = out.get("overview") or {}
    # The page mixes two sources on its own Overview: deals and baseline doors/deal come
    # from capture-data (base_supply / deals), but "SFR supply" is county-data's
    # baseline.list_size. The two supplies differ in some counties (Spotsylvania 47,615 vs
    # 49,628), so each figure is checked against the source the page actually used.
    for label, sv in [("Investor purchases (6mo)", cb["deals"]),
                      ("Baseline doors/deal", cb["dpd"]),
                      ("SFR supply", (c.get("baseline") or {}).get("list_size"))]:
        wv = _num(ov.get(label))
        if wv is not None and sv is not None and abs(float(wv) - float(sv)) > 0.05:
            check["mismatches"].append(f"{label}: workbook {wv} vs shard {sv}")

    # Which ladder did the export carry? The page's default chip is `ai`; match on the
    # rung label sequence rather than assume it.
    variants = cap.get("variants") or {}
    wb_labels = [r["signal"] for r in out.get("rows") or []]
    exported = next((k for k in ("ai", "lists")
                     if [v.get("label") for v in variants.get(k) or []] == wb_labels), None)
    check["exported_variant"] = exported
    if exported is None:
        check["mismatches"].append("workbook rungs match neither capture ladder variant")
    else:
        for row, v in zip(out["rows"], variants[exported]):
            row["list_size"] = v.get("list_size")
            row["keys"] = v.get("keys")
            row["est"] = v.get("est")
            row["thin"] = v.get("thin")
            if row["deals"] is not None and v.get("marg_deals") is not None \
                    and int(row["deals"]) != int(v["marg_deals"]):
                check["mismatches"].append(
                    f"rung {row['rung']} {row['signal']}: workbook {row['deals']} new deals "
                    f"vs shard {v['marg_deals']}")
    check["workbook_rungs"] = len(wb_labels)
    if check["mismatches"]:
        check["status"] = "MISMATCH"
    out["shard_check"] = check
    out["rows_no_ai"] = rows_from_capture(variants.get("lists"), cb["dpd"])
    out["shard_extra"] = _shard_extra(c, cap)


def from_shard_only(shards: dict, fips: str) -> dict:
    c = (shards.get("county-data") or {}).get(fips) or {}
    cap = (shards.get("capture-data") or {}).get(fips)
    if not cap:
        raise SystemExit(f"FIPS {fips} not in the capture-data shard (no Work Order for it)")
    f = ((shards.get("ftm-data") or {}).get(fips) or {}).get("rows") or []
    band_map = {"A": "Pull first (A)", "B": "Pull next (B)", "C": "Pull last (C)"}
    cb = _capture_baseline(cap)
    variants = cap.get("variants") or {}
    mo = c.get("margin_overall") or {}
    return {
        "source_xlsx": None, "format": "work_order",
        "title": f"{c.get('name') or cap.get('name')}, {c.get('state') or cap.get('st')} (FIPS {fips})",
        "overview": {"Investor purchases (6mo)": cb["deals"],
                     "Baseline doors/deal": cb["dpd"],
                     "SFR supply": (c.get("baseline") or {}).get("list_size"),
                     "Typical built-in equity": (c.get("baseline") or {}).get("median_gross"),
                     "Typical margin": mo.get("median_margin_pct"),
                     "Median purchase price": mo.get("median_purchase"),
                     "Institutional share": c.get("portfolio_pct"),
                     "Foreclosure regime": "judicial" if c.get("judicial") else "non-judicial",
                     "Data window": cap.get("window") or c.get("window")},
        "baseline": cb,
        "rows": rows_from_capture(variants.get("ai"), cb["dpd"]),
        "rows_no_ai": rows_from_capture(variants.get("lists"), cb["dpd"]),
        "ladder_measured": cap.get("built"),
        "top_zips": [{"zip": z["zip"], "stars": z.get("stars"), "deals": z.get("deals"),
                      "typical_gross": z.get("median_gross"), "dpd": z.get("dpd"),
                      "momentum": z.get("momentum")} for z in (c.get("zips") or [])],
        "price_bands": c.get("price_bands"),
        "ftm": [{"priority": band_map.get(r.get("p"), r.get("p")), "data_type": r.get("t"),
                 "office_official": r.get("o"), "phone": r.get("ph"), "source_url": r.get("u"),
                 "office_address": r.get("a"), "records_foia": r.get("foia"),
                 "access": r.get("d"), "updates": r.get("c"),
                 "verified": "Verified" if r.get("v") else "Verify locally", "notes": r.get("n")}
                for r in f],
        "ftm_narrative": [],
        "siftmap_gaps": {"from_shard_ftm_verdicts": c.get("ftm")},
        "shard_extra": _shard_extra(c, cap),
    }


def main() -> int:
    ap = argparse.ArgumentParser(description="Extract a county's doors-per-deal playbook data")
    ap.add_argument("--fips", required=True, help="5-digit county FIPS, e.g. 11001")
    ap.add_argument("--xlsx", help="the single-county Work Order workbook exported from the playbook page")
    ap.add_argument("--fetch", action="store_true",
                    help="download the playbook's free JSON shards (cross-check + list sizes, keys "
                         "and the no-AI ladder; or the whole source without --xlsx)")
    ap.add_argument("--out")
    a = ap.parse_args()

    j = BY_FIPS.get(a.fips)
    if not j:
        raise SystemExit(f"FIPS {a.fips} is not one of the 14 target jurisdictions "
                         "(src/dpd/jurisdictions.py)")
    out_path = Path(a.out) if a.out else ROOT / "data" / f"dpd_playbook_{a.fips}.json"
    shard_path = ROOT / "data" / f"dpd_playbook_{a.fips}_shards.json"

    shards = None
    if a.fetch:
        shards = fetch_shards(a.fips[:2], shard_path)
        print(f"fetched {', '.join(SHARDS)} for state {a.fips[:2]} -> {shard_path.name}")

    if a.xlsx:
        out = parse_xlsx(Path(a.xlsx))
        if out.get("fips_in_title") and out["fips_in_title"] != a.fips:
            raise SystemExit(f"workbook is for FIPS {out['fips_in_title']}, not {a.fips}")
        if shards:
            merge_shard(out, shards, a.fips)
    elif shards:
        out = from_shard_only(shards, a.fips)
    else:
        raise SystemExit("give --xlsx and/or --fetch")

    out.update({"fips": a.fips, "jurisdiction": j.name, "state_abbr": j.state_abbr,
                "extracted_at": datetime.now().isoformat(timespec="seconds")})
    out_path.write_text(json.dumps(out, indent=1), encoding="utf-8")

    def counts(rows):
        by_p: dict = {}
        for r in rows:
            by_p[r["priority"]] = by_p.get(r["priority"], 0) + 1
        return f"{{P1 {by_p.get(1, 0)}, P2 {by_p.get(2, 0)}, P3 {by_p.get(3, 0)}}}"

    rows = out["rows"]
    b = out.get("baseline") or {}
    print(f"{j.name} ({a.fips}): baseline {b.get('dpd')} doors/deal on {b.get('deals')} deals; "
          f"{len(rows)} rungs {counts(rows)}, {sum(1 for r in rows if r.get('ai'))} AI; "
          f"{len(out.get('top_zips') or [])} rated ZIPs, {len(out.get('ftm') or [])} FTM rows")
    if out.get("rows_no_ai") is not None:
        print(f"  no-AI ladder: {len(out['rows_no_ai'])} rungs {counts(out['rows_no_ai'])}")
    sc = out.get("shard_check")
    if sc:
        print(f"shard cross-check: {sc['status']}  exported ladder {sc.get('exported_variant')}  "
              f"rungs {sc.get('workbook_rungs')}")
        for m in sc.get("mismatches") or []:
            print("  MISMATCH", m)
    print(f"wrote {out_path}")
    return 1 if (sc and sc["status"] != "ok") else 0


if __name__ == "__main__":
    raise SystemExit(main())
