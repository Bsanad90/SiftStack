"""Relative and PR name repair -- fill the empty REL{n}: Full Name / PR First|Last Name
fields behind Rel{n}.{m} / Pr.{n} tagged phones from the vendor skip-trace files.

The dialers' report (via Basem, 2026-09-15): numbers on a record carry tags such as
`Rel1.1` or `Pr.2` but the REL / PR name fields are empty, so nobody knows who they
are dialing. A 150-record live sample put it at ~4% of Rel-tagged and ~24% of Pr-tagged
records (~500 / ~1,500 account-wide). The names exist in the vendor files on Drive
(`*_ReadyForDialing.xlsx`, `*_SkipTraced.xlsx`; fetched by vendor_lists_fetch.py) beside
the very same phone numbers, so the join is property address + phone digits.

Stages (each writes under output/rel_pr_name_repair/):

    pull    page every record whose phones carry a Rel/Pr/PR/Owner/Plaintiff/Defendant/
            Deceased/Spouse tag (query key `any_phone_tags`, phone-tag UUIDs from
            /api/internal/phone/tag/), then hydrate detail + custom fields. Threaded,
            resumable (records.json checkpoint). A tag past the 10k offset ceiling is
            split on `created` ranges the way live_pull does.
    index   parse every downloaded vendor file into persons {role, name, phones,
            property address} -> index.json (by phone digits).
    plan    join records to the index; decide per record which REL / PR fields to write
            and which phone tags are mislabelled (a Rel-tagged number that is the PR's,
            the 859 Lake Gaston case) -> plan.json + plan.csv.
    probe   dry-run + commit ONE record by key, then read it back.
    commit  write the plan (backup of prior values first), read back every write.
    qa      re-read every planned record against the plan, report failing keys.
    undo    restore prior field values and phone tags from a backup file.

Rules carried from earlier repairs (CLAUDE.md): a REL field written by another plan is
never blanked; the PR fields are written only when empty or when they hold a value from
the known junk families; phone tag changes go through upsert-phones with the FULL tag
list and are logged number by number; every write is read back.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import threading
import time
import warnings
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT / "src"), str(ROOT / "src" / "scripts")]

import openpyxl  # noqa: E402

from dpd_lane_assignee_report import PAGE, count  # noqa: E402
from dpd_lane_phone_score import Api, jload, jsave, say, tag_names  # noqa: E402
from live_pull import EARLIEST_GUESS, _iso  # noqa: E402
from obituary_dp_batch import name_tokens, names_match, sift_key, zip5  # noqa: E402
from pr_repair_copr_multi import street_matches  # noqa: E402

RUN = ROOT / "output" / "rel_pr_name_repair"
RECORDS = RUN / "records.json"
COHORT = RUN / "cohort_uuids.json"
PHONE_TAGS = RUN / "phone_tags.json"
INDEX = RUN / "index.json"
PLAN = RUN / "plan.json"
VENDOR_FILES = ROOT / "output" / "vendor_lists" / "files"

OFFSET_CEILING = 10000
SAFE_LIMIT = 9000

# Phone-tag families that name WHOSE number it is. Both spellings live on the account
# ("Rel1.1" and "Rel 1.1", "Pr.1"/"PR.1"/"Pr4", "Owner.1"/"Owner 1").
FAMILY_RE = re.compile(
    r"^(?P<kind>rel|pr|owner|plaintiff|defendant|deceased|spouse|attorney)"
    r"\.? ?(?P<n>\d+)?(?:[. ](?P<m>\d+))?$", re.I)


def d10(n) -> str:
    """Last 10 digits. A float cell (4403412393.0, how openpyxl hands back most vendor
    phone cells) must lose its '.0' FIRST -- str() then digits would shift it to
    '4034123930' and index the number under a garbage key (~70k keys in the first build)."""
    if isinstance(n, float):
        n = int(n) if n.is_integer() else n
    s = str(n or "").strip()
    s = re.sub(r"\.0+$", "", s)
    return "".join(c for c in s if c.isdigit())[-10:]


def parse_tag(t: str):
    """'Rel3.2' -> ('rel', 3, 2); 'Pr.1' -> ('pr', 1, None); 'PR2.1' -> ('pr', 2, 1)."""
    m = FAMILY_RE.match((t or "").strip())
    if not m:
        return None
    return (m.group("kind").lower(),
            int(m.group("n")) if m.group("n") else None,
            int(m.group("m")) if m.group("m") else None)


# ----------------------------------------------------------------------------- pull

def fetch_phone_tags(api: Api) -> list[dict]:
    out, off = [], 0
    while True:
        st, r = api.get(f"/api/internal/phone/tag/?limit=200&offset={off}")
        if st != 200:
            raise RuntimeError(f"phone tag list {st}: {str(r)[:200]}")
        out += r.get("results") or []
        off += 200
        if not r.get("next"):
            break
    return out


def page_uuids(api: Api, query: dict, expected: int) -> list[str]:
    out, offset = [], 0
    while offset < expected and offset < OFFSET_CEILING:
        st, r = api.post_as_get("/api/internal/property/",
                                {"limit": PAGE, "offset": offset, "query": query})
        if st != 200:
            raise RuntimeError(f"list failed at offset {offset}: {st} {str(r)[:300]}")
        rows = r.get("results") or []
        if not rows:
            break
        out += [row["uuid"] for row in rows]
        offset += PAGE
    return out


def page_tag(api: Api, tag_uuid: str, n: int) -> list[str]:
    """All uuids carrying one phone tag. Past the offset ceiling the query is split on
    `created` ranges (live_pull's proven shape) until every bucket is under SAFE_LIMIT."""
    base = {"any_phone_tags": [tag_uuid]}
    if n < SAFE_LIMIT:
        return page_uuids(api, {"must": base}, n)

    def rng(start, end):
        return {"must": {**base, "created": [_iso(start), _iso(end)]}}

    def split(start, end):
        c = count(api, rng(start, end))
        if c == 0:
            return []
        if c <= SAFE_LIMIT:
            return [(start, end, c)]
        if (end - start) <= timedelta(hours=1):
            raise RuntimeError(f"{c} records created within one hour carry this tag -- "
                               "cannot page under the offset ceiling")
        mid = start + (end - start) / 2
        return split(start, mid) + split(mid, end)

    buckets = split(EARLIEST_GUESS, datetime.now(timezone.utc) + timedelta(days=1))
    out: list[str] = []
    for s, e, c in buckets:
        out += page_uuids(api, rng(s, e), c)
    say(f"    split into {len(buckets)} created-range buckets -> {len(set(out))} uuids (count said {n})")
    return list(dict.fromkeys(out))


def build_cohort(api: Api) -> dict[str, dict]:
    """{uuid: {tags: [phone-tag titles that selected it]}} for every family tag.

    Per-tag paging keeps each query under the 10,000 offset ceiling; Rel3.1 alone is
    10,446 records, hence the created-range split. The account total is the guard
    against an ignored filter key.
    """
    tags = fetch_phone_tags(api)
    jsave(PHONE_TAGS, tags)
    fam = [t for t in tags if parse_tag(t["title"])]
    say(f"{len(tags)} phone tags on the account, {len(fam)} in the who-is-this families")
    total = count(api, {})
    cohort: dict[str, dict] = jload(COHORT, {}) if COHORT.exists() else {}
    done_tags = set(cohort.pop("_done_tags", []))
    for t in fam:
        if t["title"] in done_tags:
            continue
        q = {"must": {"any_phone_tags": [t["uuid"]]}}
        n = count(api, q)
        if n == total:
            raise RuntimeError("filtered count equals the account total -- the filter was ignored")
        uu = page_tag(api, t["uuid"], n) if n else []
        for u in uu:
            cohort.setdefault(u, {"tags": []})["tags"].append(t["title"])
        done_tags.add(t["title"])
        say(f"  {t['title']:<14} {n:>6} -> cohort {len(cohort)}")
        jsave(COHORT, {**cohort, "_done_tags": sorted(done_tags)})
    jsave(COHORT, {**cohort, "_done_tags": sorted(done_tags)})
    return cohort


def cf_values(cf) -> dict[str, str]:
    """{label: value} from the per-record custom-field endpoint, ACTIVE fields only.

    Rows carry `custom_field.is_active` and `uuid`; the account has an inactive twin of
    'REL1: Full Name', so a label-keyed last-wins dict could let a dead row blank the live
    one. Verified 2026-09-15: the endpoint returns only rows with values, active so far."""
    rows = cf if isinstance(cf, list) else ((cf.get("results") or cf.get("data") or [])
                                            if isinstance(cf, dict) else [])
    out: dict[str, str] = {}
    for r in rows:
        f = r.get("custom_field") or {}
        if f.get("is_active") is False:
            continue
        lbl = f.get("label") or ""
        if lbl and (r.get("value") or "" or lbl not in out):
            out[lbl] = r.get("value") or ""
    return out


def hydrate_one(api: Api, u: str) -> dict:
    st, body = api.get(f"/api/internal/property/{u}/")
    if st == 404:
        return {"_gone": True}
    if st != 200:
        return {"_error": st}
    stc, cf = api.get(f"/api/internal/property/{u}/custom-field/")
    if stc != 200:
        # a failed field read must never look like "all fields empty" (the plan would
        # then overwrite another plan's REL name)
        return {"_error": f"cf {stc}"}
    cfv = cf_values(cf)
    addr = body.get("address") or {}
    owner = body.get("owner") or {}
    oaddr = owner.get("address") or {}
    return {
        "street": addr.get("street"), "city": addr.get("city"), "state": addr.get("state"),
        "zip": addr.get("zip5") or (addr.get("postal_code") or "")[:5], "county": addr.get("county"),
        "status": body.get("status"), "type": body.get("type"),
        "owner_uuid": owner.get("uuid"),
        "owner_first": owner.get("first_name"), "owner_last": owner.get("last_name"),
        "owner_company": owner.get("company"), "owner_deceased": owner.get("deceased"),
        "mail_street": oaddr.get("street"), "mail_city": oaddr.get("city"),
        "mail_state": oaddr.get("state"), "mail_zip": oaddr.get("zip5") or (oaddr.get("postal_code") or "")[:5],
        "lists": tag_names(body.get("lists")), "tags": tag_names(body.get("tags")),
        "assigned_to": body.get("assigned_to"),
        "personal_representative": body.get("personal_representative"),
        "phones": [{"number": p.get("number"), "type": p.get("type"),
                    "tags": tag_names(p.get("tags")), "status": p.get("status"),
                    "is_connected": p.get("is_connected", p.get("isConnected")),
                    "verified": p.get("verified")}
                   for p in owner.get("phones") or []],
        "cf": {k: v for k, v in cfv.items() if v},
        "cf_status": stc,
        "pulled": datetime.now().isoformat(timespec="seconds"),
    }


def cmd_pull(a) -> int:
    RUN.mkdir(parents=True, exist_ok=True)
    api = Api()
    say("JWT minted OK")
    cohort = build_cohort(api)
    uuids = [u for u in cohort if not u.startswith("_")]
    say(f"cohort: {len(uuids)} records")

    records: dict = {} if a.fresh else jload(RECORDS, {})
    todo = [u for u in uuids if u not in records or "_error" in records[u]
            or records[u].get("cf_status") not in (None, 200)]
    if a.limit:
        todo = todo[:a.limit]
    say(f"hydrating {len(todo)} ({len(records)} cached) with {a.threads} threads...")
    apis = [api] + [Api() for _ in range(a.threads - 1)]
    lock = threading.Lock()
    t0 = time.time()
    done = 0

    def work(i_u):
        i, u = i_u
        return u, hydrate_one(apis[i % len(apis)], u)

    with ThreadPoolExecutor(max_workers=a.threads) as ex:
        futs = [ex.submit(work, (i, u)) for i, u in enumerate(todo)]
        for f in as_completed(futs):
            u, rec = f.result()
            with lock:
                rec["selected_by"] = cohort[u]["tags"]
                records[u] = rec
                done += 1
                if done % 200 == 0 or done == len(todo):
                    jsave(RECORDS, records)
                    rate = done / max(time.time() - t0, 1)
                    say(f"  {done}/{len(todo)}  {rate:.2f} rec/s  eta {((len(todo) - done) / max(rate, 0.01)) / 60:.0f} min")
    jsave(RECORDS, records)
    gone = sum(1 for r in records.values() if r.get("_gone"))
    err = sum(1 for r in records.values() if r.get("_error"))
    say(f"done: {len(records)} records ({gone} gone, {err} errors)")
    return 0


# ----------------------------------------------------------------------------- index

PROP_STREET_COLS = ["property address", "property street"]
PROP_ZIP_COLS = ["property zip", "property zip code", "property zipcode"]
INPUT_STREET_COLS = ["input: address 1", "owner street", "owner address", "mailing address", "pr address",
                     "pr address 1"]
INPUT_ZIP_COLS = ["input: zip code", "owner zip code", "owner zipcode", "mailing zip", "pr zip code"]
# PR skip files carry the PROPERTY in the vendor's pass-through columns: INPUT: Extra 1-4 /
# Other 1-4 = street, city, state, zip (12 of 26 sampled files). Never pair the INPUT street
# (the PR's mailing) with one of these zips.
PASSTHRU_SETS = [("input: extra 1", "input: extra 4"), ("input: other 1", "input: other 4")]
STREET_LIKE_RE = re.compile(r"^\d+\s+\S+")

PH_RE = re.compile(r"^ph: ?phone ?(\d+)$", re.I)
PLAIN_PHONE_RE = re.compile(r"^phone ?(\d+)$", re.I)
PLAIN_TAG_RE = re.compile(r"^phone ?(\d+) tags$", re.I)
REL_NAME_RE = re.compile(r"^rel ?(\d+): ?full name$", re.I)
REL_FIRST_RE = re.compile(r"^rel ?(\d+): ?first name$", re.I)
REL_LAST_RE = re.compile(r"^rel ?(\d+): ?last name$", re.I)
REL_PHONE_RE = re.compile(r"^rel ?(\d+): ?phone ?(\d+)$", re.I)

# The 2026-09-04 PR repair's junk families: match-notes, sheet labels, row numbers,
# estate numbers, dates, booleans -- and identical first/last (checked separately).
JUNK_PR_RE = re.compile(r"^(matched|legal notices|\d+(\.0)?$|false$|true$|\d{4}-\d{2}-\d{2}|n/?a$|none$|null$|#n/a)",
                        re.I)


def clean_name(s) -> str:
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    if s.lower() in ("none", "null", "nan", "n/a", "#n/a"):
        return ""
    return s


def title_name(s: str) -> str:
    """Vendor names are shouted ('JAMES RAPER COBB '); the account reads better in title
    case. Keeps Mc/O' and hyphens sensible; roman suffixes stay upper."""
    out = []
    for w in clean_name(s).split(" "):
        lw = w.lower()
        if lw in ("ii", "iii", "iv"):
            out.append(w.upper())
        else:
            part = "-".join(x[:1].upper() + x[1:].lower() for x in w.split("-"))
            part = re.sub(r"^(Mc)(\w)", lambda m: m.group(1) + m.group(2).upper(), part)
            part = re.sub(r"^(O')(\w)", lambda m: m.group(1) + m.group(2).upper(), part)
            out.append(part)
    return " ".join(out)


def name_key(s: str) -> str:
    """Grouping key: 'JOSEFF W COCKRILL' and 'Joseff Cockrill' are one person (middle
    initials dropped, suffixes/ordinals already stripped by name_tokens)."""
    toks = name_tokens(s)
    return " ".join(t for t in toks if len(t) > 1) or " ".join(toks)


def _get(row: dict, cols: list[str]) -> str:
    for c in cols:
        v = row.get(c)
        if v and v.lower() not in ("none", "null"):
            return v
    return ""


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def file_role(name: str) -> str:
    """Whose phones the INPUT / PH columns describe, read off the file name.

    'Monthly Probate_202604 SkipTraced' input person = the PR (proven: Samuel Cantanese is
    the PR in the matching ReadyForDialing). 'Probate No PR' / 'Decd' / 'Inheritance' /
    'Obituar' skips are of the decedent. Divorce files name their party.
    """
    n = name.lower()
    if "plaintiff" in n:
        return "plaintiff"
    if "defendant" in n:
        return "defendant"
    if re.search(r"no pr|decd|deceased|inheritance|obituar|decedent", n):
        return "deceased"
    if re.search(r"\bpr\b|probate", n):
        return "pr"
    return "owner"


def parse_sheet(ws, fname: str, sheet: str) -> list[dict]:
    rows = ws.iter_rows(values_only=True)
    hdr = next(rows, None)
    if not hdr:
        return []
    cols = [re.sub(r"\s+", " ", str(h or "").strip()) for h in hdr]
    lc = [c.lower() for c in cols]
    persons: list[dict] = []
    role_of_file = file_role(fname)

    has_pr_cols = "pr first name" in lc or "pr last name" in lc
    has_deceased = "deceased first name" in lc
    has_input = "input: first name" in lc
    has_generic = "first name" in lc and "last name" in lc
    has_party = "plaintiff first name" in lc or "defendant first name" in lc
    rel_name_idx = {int(REL_NAME_RE.match(c).group(1)): i for i, c in enumerate(lc) if REL_NAME_RE.match(c)}
    rel_first_idx = {int(REL_FIRST_RE.match(c).group(1)): i for i, c in enumerate(lc) if REL_FIRST_RE.match(c)}
    rel_last_idx = {int(REL_LAST_RE.match(c).group(1)): i for i, c in enumerate(lc) if REL_LAST_RE.match(c)}
    rel_phone_idx: dict[int, list[int]] = {}
    for i, c in enumerate(lc):
        m = REL_PHONE_RE.match(c)
        if m:
            rel_phone_idx.setdefault(int(m.group(1)), []).append(i)
    ph_idx = [i for i, c in enumerate(lc) if PH_RE.match(c)]
    plain_idx = [(i, int(PLAIN_PHONE_RE.match(c).group(1))) for i, c in enumerate(lc) if PLAIN_PHONE_RE.match(c)]
    plain_tag_idx = {int(PLAIN_TAG_RE.match(c).group(1)): i for i, c in enumerate(lc) if PLAIN_TAG_RE.match(c)}

    def val(raw, i):
        return _cell(raw[i]) if i is not None and i < len(raw) else ""

    def plain_phones(raw, want_prefix: str) -> list[str]:
        """'Phone N' cells credited to a party only by their 'Phone N Tags' cell; when the
        sheet has tag columns an UNTAGGED number is nobody's (it could be a relative's)."""
        out = []
        for i, k in plain_idx:
            tg = val(raw, plain_tag_idx.get(k))
            if len(d10(raw[i] if i < len(raw) else "")) != 10:
                continue
            if plain_tag_idx and not tg:
                continue
            if not tg or tg.lower().startswith(want_prefix):
                out.append(d10(raw[i]))
        return out

    for rn, raw in enumerate(rows, 2):
        if raw is None or all(v is None for v in raw):
            continue
        row = {lc[i]: val(raw, i) for i in range(len(lc))}
        prop_street = _get(row, PROP_STREET_COLS)
        prop_zip = _get(row, PROP_ZIP_COLS)
        in_street = _get(row, INPUT_STREET_COLS)
        in_zip = _get(row, INPUT_ZIP_COLS)
        if not prop_street:
            for s_col, z_col in PASSTHRU_SETS:
                cand = row.get(s_col, "")
                if cand and STREET_LIKE_RE.match(cand):
                    prop_street, prop_zip = cand, (row.get(z_col, "") if re.fullmatch(r"\d{5}(-\d{4})?", row.get(z_col, "") or "") else "")
                    break
        if not prop_street and has_input and not has_pr_cols and role_of_file == "deceased":
            # a decedent skip's INPUT address is the property itself
            prop_street, prop_zip = in_street, in_zip
        base = {"file": fname, "sheet": sheet, "row": rn,
                "prop_key": sift_key(prop_street, prop_zip) if prop_street else "",
                "prop_street": prop_street, "prop_zip": zip5(prop_zip),
                "input_key": sift_key(in_street, in_zip) if in_street else "",
                "input_street": in_street}
        anchor = ""
        if has_deceased:
            anchor = clean_name(f"{row.get('deceased first name', '')} {row.get('deceased last name', '')}")
        elif has_input:
            anchor = clean_name(f"{row.get('input: first name', '')} {row.get('input: last name', '')}")
        elif has_generic and not has_pr_cols:
            anchor = clean_name(f"{row.get('first name', '')} {row.get('last name', '')}")
        base["anchor"] = anchor

        ph_phones = [d10(raw[i]) for i in ph_idx if i < len(raw) and len(d10(raw[i])) == 10]

        # ---- the PH: / Phone N columns belong to one person per row
        if has_pr_cols:
            first, last = clean_name(row.get("pr first name")), clean_name(row.get("pr last name"))
            mid = clean_name(row.get("pr middle name"))
            phones = ph_phones or plain_phones(raw, "pr")
            if (first or last) and phones:
                persons.append({**base, "role": "pr", "first": first, "last": last,
                                "name": clean_name(f"{first} {mid} {last}"), "phones": phones})
        elif has_input:
            first, last = clean_name(row.get("input: first name")), clean_name(row.get("input: last name"))
            if (first or last) and ph_phones:
                persons.append({**base, "role": role_of_file, "first": first, "last": last,
                                "name": clean_name(f"{first} {last}"), "phones": ph_phones})
        elif has_party:
            for party in ("plaintiff", "defendant"):
                first, last = clean_name(row.get(f"{party} first name")), clean_name(row.get(f"{party} last name"))
                if not (first or last):
                    continue
                phones = plain_phones(raw, party) if plain_idx else (ph_phones if role_of_file == party else [])
                if phones:
                    persons.append({**base, "role": party, "first": first, "last": last,
                                    "name": clean_name(f"{first} {last}"), "phones": phones})
        elif has_generic:
            first, last = clean_name(row.get("first name")), clean_name(row.get("last name"))
            if (first or last) and ph_phones:
                persons.append({**base, "role": "owner" if role_of_file == "pr" else role_of_file,
                                "first": first, "last": last, "name": clean_name(f"{first} {last}"),
                                "phones": ph_phones})
        elif has_deceased and ph_phones:
            first, last = clean_name(row.get("deceased first name")), clean_name(row.get("deceased last name"))
            persons.append({**base, "role": "deceased", "first": first, "last": last,
                            "name": clean_name(f"{first} {last}"), "phones": ph_phones})

        # ---- relatives
        for n in sorted(set(rel_name_idx) | set(rel_first_idx)):
            nm = clean_name(val(raw, rel_name_idx.get(n)))
            first = clean_name(val(raw, rel_first_idx.get(n)))
            last = clean_name(val(raw, rel_last_idx.get(n)))
            if not nm:
                nm = clean_name(f"{first} {last}")
            if not first and nm:
                toks = nm.split()
                first, last = toks[0], (toks[-1] if len(toks) > 1 else "")
            phones = [d10(raw[i]) for i in rel_phone_idx.get(n, []) if i < len(raw) and len(d10(raw[i])) == 10]
            if nm and phones:
                persons.append({**base, "role": "rel", "rel_n": n, "first": first, "last": last,
                                "name": nm, "phones": phones})
    return persons


def cmd_index(a) -> int:
    RUN.mkdir(parents=True, exist_ok=True)
    warnings.filterwarnings("ignore")
    files = sorted(VENDOR_FILES.rglob("*.xls*")) + sorted(VENDOR_FILES.rglob("*.csv"))
    say(f"{len(files)} vendor files")
    if not files:
        say("FAIL: no vendor files under output/vendor_lists/files -- run vendor_lists_fetch.py first")
        return 2
    persons: list[dict] = []
    fails = []
    for i, f in enumerate(files, 1):
        rel = str(f.relative_to(VENDOR_FILES))
        try:
            if f.suffix.lower() == ".csv":
                wb = openpyxl.Workbook()
                ws = wb.active
                with open(f, encoding="utf-8-sig", errors="ignore", newline="") as fh:
                    for r in csv.reader(fh):
                        ws.append(r)
                persons += parse_sheet(ws, rel, "csv")
            else:
                wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
                for ws in wb.worksheets:
                    persons += parse_sheet(ws, rel, ws.title)
        except Exception as exc:  # noqa: BLE001
            fails.append((rel, str(exc)[:120]))
        if i % 100 == 0:
            say(f"  {i}/{len(files)} files, {len(persons)} persons")
    by_phone: dict[str, list[int]] = {}
    for pid, p in enumerate(persons):
        p["id"] = pid
        for ph in set(p["phones"]):
            by_phone.setdefault(ph, []).append(pid)
    roles: dict[str, int] = {}
    for p in persons:
        roles[p["role"]] = roles.get(p["role"], 0) + 1
    jsave(INDEX, {"built": datetime.now().isoformat(timespec="seconds"), "files": len(files),
                  "failed_files": fails, "persons": persons, "by_phone": by_phone})
    say(f"index: {len(persons)} persons, {len(by_phone)} distinct phones, roles {roles}, "
        f"{len(fails)} unreadable files")
    for f_, e in fails[:10]:
        say(f"  FAIL {f_}: {e}")
    return 0


# ----------------------------------------------------------------------------- plan

CONF_ORDER = {"HIGH": 0, "MEDIUM": 1, "LOW": 2}


def record_keys(r: dict) -> tuple[str, str]:
    return (sift_key(r.get("street"), r.get("zip")),
            sift_key(r.get("mail_street"), r.get("mail_zip")) if r.get("mail_street") else "")


def evidence(person: dict, r: dict, prop_key: str, mail_key: str) -> tuple[str, list[str]]:
    """How sure we are this vendor row is about THIS record."""
    ev = []
    if person["prop_key"] and person["prop_key"] == prop_key:
        ev.append("property address")
    elif person["prop_street"] and r.get("street") and street_matches(person["prop_street"], r["street"]) \
            and (not person["prop_zip"] or not r.get("zip") or person["prop_zip"] == r["zip"]):
        ev.append("property street")
    if person["input_key"] and person["input_key"] in (prop_key, mail_key):
        ev.append("input address")
    owner = clean_name(f"{r.get('owner_first') or ''} {r.get('owner_last') or ''}")
    if person.get("anchor") and owner and names_match(person["anchor"], owner):
        ev.append("owner name")
    if person["role"] == "pr":
        pr_now = clean_name(f"{r['cf'].get('PR First Name', '')} {r['cf'].get('PR Last Name', '')}")
        if pr_now and names_match(person["name"], pr_now):
            ev.append("PR name")
    if any(e.startswith("property") for e in ev):
        return "HIGH", ev
    if ev:
        return "MEDIUM", ev
    return "LOW", ev


def is_junk_pr(first, last) -> bool:
    """Empty, or one of the 2026-09-04 junk families. Identical first/last was that
    repair's signature, but a real 'Thomas Thomas' exists: identical counts as junk only
    when the value is not a plain alphabetic name."""
    f, l = clean_name(first), clean_name(last)
    if not f and not l:
        return True
    if JUNK_PR_RE.match(f or l):
        return True
    if f and l and f.lower() == l.lower() and not re.fullmatch(r"[A-Za-z' .-]+", f):
        return True
    return False


def resolve_group(hits: list[dict], r: dict, prop_key: str, mail_key: str, want_role: str,
                  slot_phones: set[str], slot_n: int | None = None, use_rel_n: bool = True) -> dict:
    """Pick ONE vendor person for a set of phones. Groups the hits by person (name_key,
    then merged by names_match so 'WARREN J SIEMINSKI COCKRILL' and 'Warren Cockrill' are
    one) and ranks the groups by: best address/name evidence, how many of the slot's
    numbers the person holds, whether the vendor row put them in the SAME REL column
    (only when `use_rel_n` -- the caller disables it on records whose slots drifted from
    the vendor columns), and whether their role is the wanted one. A tie on all four
    between different names is a review, never a guess.

    A vendor row that lists its own decedent / input person as a relative (34 rows in the
    April-2024 pre-probate list) is dropped before grouping."""
    groups: dict[str, dict] = {}
    keys: list[str] = []
    for p in hits:
        if p.get("anchor") and names_match(p["name"], p["anchor"]) and p["role"] == "rel":
            continue
        conf, ev = evidence(p, r, prop_key, mail_key)
        k = name_key(p["name"]) or p["name"].lower()
        if k not in groups:
            k = next((k2 for k2 in keys if names_match(groups[k2]["name"], p["name"])), k)
        if k not in groups:
            keys.append(k)
        g = groups.setdefault(k, {"name": p["name"], "first": p["first"], "last": p["last"],
                                  "roles": set(), "conf": "LOW", "evidence": set(), "files": set(),
                                  "rel_n": set(), "phones": set()})
        g["roles"].add(p["role"])
        g["files"].add(p["file"])
        g["phones"].update(p["phones"])
        if p.get("rel_n"):
            g["rel_n"].add(p["rel_n"])
        g["evidence"].update(ev)
        if CONF_ORDER[conf] < CONF_ORDER[g["conf"]] or (
                CONF_ORDER[conf] == CONF_ORDER[g["conf"]] and len(p["name"]) > len(g["name"])):
            g["conf"] = conf
            g["name"], g["first"], g["last"] = p["name"], p["first"], p["last"]
    if not groups:
        return {"verdict": "no_vendor_hit"}

    def rank(g):
        # holding EVERY number in the slot comes before address confidence: a PR whose
        # property-matched row holds one of a relative's two numbers must not outrank the
        # relative who holds both (859 Lake Gaston REL5, Warren vs the PR)
        return (0 if slot_phones <= g["phones"] else 1,
                CONF_ORDER[g["conf"]],
                -len(g["phones"] & slot_phones),
                0 if (use_rel_n and slot_n and slot_n in g["rel_n"]) else 1,
                0 if want_role in g["roles"] else 1,
                -len(g["evidence"]))

    ranked = sorted(groups.values(), key=rank)
    best = ranked[0]
    if len(ranked) > 1 and rank(ranked[1])[:5] == rank(best)[:5]:
        return {"verdict": "ambiguous", "names": [g["name"] for g in ranked[:4]], "conf": best["conf"],
                "tie_on_rel_n": (not use_rel_n) and any(slot_n and slot_n in g["rel_n"] for g in ranked[:2])}
    roles = best["roles"]
    out = {"verdict": "ok", "name": best["name"], "first": best["first"], "last": best["last"],
           "conf": best["conf"], "evidence": sorted(best["evidence"]), "files": sorted(best["files"])[:3],
           "roles": sorted(roles), "vendor_rel_n": sorted(best["rel_n"]),
           "held_phones": sorted(best["phones"] & slot_phones),
           "coverage": f"{len(best['phones'] & slot_phones)}/{len(slot_phones)}"}
    if want_role == "rel" and "rel" not in roles:
        # a Rel-tagged number that the vendor says is the PR's (859 Lake Gaston), or the
        # decedent's / a divorce party's
        out["verdict"] = "mislabelled_pr" if roles == {"pr"} else "role_mismatch"
    elif want_role == "pr" and "pr" not in roles:
        out["verdict"] = "role_mismatch"
    return out


def plan_record(u: str, r: dict, idx: dict) -> dict | None:
    by_phone, persons = idx["by_phone"], idx["persons"]
    cf = r.get("cf") or {}
    prop_key, mail_key = record_keys(r)
    slots: dict[int, list[dict]] = {}
    pr_phones: list[dict] = []
    for ph in r.get("phones") or []:
        for t in ph.get("tags") or []:
            pt = parse_tag(t)
            if not pt:
                continue
            kind, n, m = pt
            if kind == "rel" and n:
                slots.setdefault(n, []).append({**ph, "tag": t, "m": m})
            elif kind == "pr":
                pr_phones.append({**ph, "tag": t, "n": n, "m": m})
    if not slots and not pr_phones:
        return None
    writes: dict[str, str] = {}
    tag_changes: list[dict] = []
    findings: list[dict] = []
    review: list[str] = []
    missing_rel = [n for n in sorted(slots) if not clean_name(cf.get(f"REL{n}: Full Name"))]
    pr_first, pr_last = cf.get("PR First Name", ""), cf.get("PR Last Name", "")
    # The PR gets filled whenever it is empty/junk and a vendor row names it -- also on a
    # record with NO Pr-tagged phone whose REL slot turns out to hold the PR's numbers.
    pr_missing = is_junk_pr(pr_first, pr_last)
    pr_missing_before = bool(pr_phones) and pr_missing
    if not missing_rel and not (pr_phones and pr_missing):
        return {"uuid": u, "street": r.get("street"), "zip": r.get("zip"), "status": "complete",
                "rel_slots": sorted(slots), "pr_tagged": len(pr_phones)}
    owner_name = clean_name(f"{r.get('owner_first') or ''} {r.get('owner_last') or ''}")

    def describe(res: dict) -> str:
        s = res["verdict"]
        if res.get("conf"):
            s += f" ({res['conf']})"
        if res.get("names"):
            s += f" names={res['names']}"
        if res.get("name"):
            s += f" {res['name']} roles={res.get('roles')}"
        return s

    def resolve_slot(n: int, use_rel_n: bool) -> dict:
        phones = slots[n]
        slot_phones = {d10(p["number"]) for p in phones}
        # A slot whose every number ALSO carries a Pr tag is the PR's (859 Lake Gaston after
        # the RFD re-tag: the same number is Rel3.1 AND Pr.2). The account already says whose
        # it is; the REL field just has to name the same person.
        also_pr = all(any((parse_tag(t) or ("",))[0] == "pr" for t in p["tags"]) for p in phones)
        hits = [persons[pid] for ph in phones for pid in by_phone.get(d10(ph["number"]), [])]
        res = resolve_group(hits, r, prop_key, mail_key, "pr" if also_pr else "rel", slot_phones, n, use_rel_n)
        res.update({"slot": f"REL{n}", "phones": sorted(slot_phones),
                    "tags": [p["tag"] for p in phones], "also_pr_tagged": also_pr})
        if also_pr and res["verdict"] == "role_mismatch" and "rel" in (res.get("roles") or []) \
                and n in (res.get("vendor_rel_n") or []):
            # shared line: the vendor's REL{n} IS this relative (Erika Cockrill, REL1), the
            # Pr tag on the same number is the PR's household line. Name the relative.
            res["verdict"] = "ok"
            res["note"] = "shared line with the PR; REL name from the vendor's same REL column"
        elif also_pr and res["verdict"] == "ok" and "rel" not in res["roles"]:
            # the PR won the slot outright; if the vendor's SAME REL column names a relative
            # who also holds every number, the REL field takes the relative and the PR
            # fields still take the PR
            rel_hits = [h for h in hits if h["role"] == "rel"]
            alt = resolve_group(rel_hits, r, prop_key, mail_key, "rel", slot_phones, n, True)
            if alt["verdict"] == "ok" and "rel" in alt["roles"] and n in (alt.get("vendor_rel_n") or []) \
                    and slot_phones <= set(alt.get("held_phones") or []) and alt["conf"] in ("HIGH", "MEDIUM"):
                pr_side = {k: res[k] for k in ("name", "first", "last", "roles", "conf")}
                res.update({k: alt[k] for k in ("name", "first", "last", "roles", "evidence", "files",
                                                 "vendor_rel_n", "held_phones", "coverage")})
                res["conf"] = alt["conf"]
                res["pr_for_fields"] = pr_side
                res["note"] = "shared line: REL name from the vendor's same REL column, PR from the PR row"
        if res["verdict"] in ("ok", "mislabelled_pr") and owner_name and names_match(res["name"], owner_name):
            # on probate records the owner IS the decedent -- a vendor row naming them as a
            # relative (or the PR) is the wrong person
            res["verdict"] = "names_owner"
        return res

    # Pass 1 without the vendor-REL-column tie-break, to learn whether this record's slot
    # numbers drifted from the vendor's columns (6849 Belclare: account REL2 = vendor REL1).
    first = {n: resolve_slot(n, False) for n in missing_rel}
    aligned = [n for n, res in first.items() if res["verdict"] == "ok" and n in (res.get("vendor_rel_n") or [])]
    drifted = [n for n, res in first.items() if res["verdict"] == "ok" and res.get("vendor_rel_n")
               and n not in res["vendor_rel_n"]]
    for n in missing_rel:
        res = first[n]
        if res["verdict"] == "ambiguous" and res.get("tie_on_rel_n") and aligned and not drifted:
            res = resolve_slot(n, True)
            res["note"] = "tie broken by the vendor REL column (other slots of this record align)"
        phones = slots[n]
        also_pr = res["also_pr_tagged"]
        findings.append(res)
        # REL names are stored SHOUTED on this account (the DP pusher .upper()s them, the
        # vendor uploads were upper); PR First/Last are title case ('Jennifer M' / 'Mckee').
        if also_pr and res["verdict"] == "ok" and res["conf"] in ("HIGH", "MEDIUM"):
            writes[f"REL{n}: Full Name"] = clean_name(res["name"]).upper()
            pr_side = res.get("pr_for_fields") or res
            if pr_missing and "pr" in pr_side["roles"] and pr_side["conf"] in ("HIGH", "MEDIUM"):
                writes["PR First Name"] = title_name(pr_side["first"])
                writes["PR Last Name"] = title_name(pr_side["last"])
                pr_missing = False
        elif res["verdict"] == "ok" and res["conf"] in ("HIGH", "MEDIUM"):
            writes[f"REL{n}: Full Name"] = clean_name(res["name"]).upper()
            for p in phones:
                if p.get("m") and 1 <= p["m"] <= 3 and not cf.get(f"REL{n}: Phone {p['m']}"):
                    # same format the DP pusher writes: the number as it sits on the owner
                    writes[f"REL{n}: Phone {p['m']}"] = str(p["number"])
        elif res["verdict"] == "mislabelled_pr" and res["conf"] in ("HIGH", "MEDIUM"):
            # the numbers the PR actually holds: name goes to the PR fields, each such tag
            # becomes the first FREE Pr.N on the owner (never a duplicate). A slot number the
            # PR does NOT hold is left alone and reported -- it may be a relative's.
            held = set(res.get("held_phones") or [])
            if pr_missing or names_match(res["name"], f"{pr_first} {pr_last}"):
                if pr_missing:
                    writes["PR First Name"] = title_name(res["first"])
                    writes["PR Last Name"] = title_name(res["last"])
                    pr_missing = False
                used = {(parse_tag(t) or ("", None, None))[1] for ph in r["phones"] for t in ph["tags"]
                        if (parse_tag(t) or ("",))[0] == "pr"}
                used |= {t["to_n"] for t in tag_changes}
                for p in phones:
                    if d10(p["number"]) not in held:
                        review.append(f"REL{n}: {d10(p['number'])} ({p['tag']}) is not among the PR's numbers "
                                      f"in the vendor file -- left tagged, name unknown")
                        continue
                    existing_pr = next((t for t in p["tags"] if (parse_tag(t) or ("",))[0] == "pr"), None)
                    if existing_pr:
                        # already carries a Pr tag: only the wrong Rel tag comes off
                        tag_changes.append({"number": p["number"], "from": p["tag"], "to": existing_pr,
                                            "to_n": (parse_tag(existing_pr) or ("", None))[1],
                                            "tags_before": p["tags"], "slot": f"REL{n}", "who": res["name"]})
                        continue
                    k = 1
                    while k in used:
                        k += 1
                    used.add(k)
                    tag_changes.append({"number": p["number"], "from": p["tag"], "to": f"Pr.{k}", "to_n": k,
                                        "tags_before": p["tags"], "slot": f"REL{n}", "who": res["name"]})
                    if p.get("m") and d10(cf.get(f"REL{n}: Phone {p['m']}", "")) == d10(p["number"]):
                        review.append(f"REL{n}: Phone {p['m']} field still holds {d10(p['number'])}, "
                                      f"now retagged {tag_changes[-1]['to']} -- clear by hand")
            else:
                review.append(f"REL{n}: vendor says these are PR {res['name']}'s numbers but the record's "
                              f"PR is {pr_first} {pr_last}")
        else:
            review.append(f"REL{n}: " + describe(res))
    if pr_missing and pr_phones:
        pr_set = {d10(p["number"]) for p in pr_phones}
        hits = [persons[pid] for ph in pr_phones for pid in by_phone.get(d10(ph["number"]), [])]
        res = resolve_group(hits, r, prop_key, mail_key, "pr", pr_set)
        res.update({"slot": "PR", "phones": sorted(pr_set),
                    "tags": [p["tag"] for p in pr_phones]})
        if res["verdict"] == "ok" and owner_name and names_match(res["name"], owner_name):
            res["verdict"] = "names_owner"
        findings.append(res)
        if res["verdict"] == "ok" and res["conf"] in ("HIGH", "MEDIUM"):
            writes["PR First Name"] = title_name(res["first"])
            writes["PR Last Name"] = title_name(res["last"])
            if clean_name(pr_first) or clean_name(pr_last):
                res["replaces_junk"] = f"{pr_first} | {pr_last}"
        else:
            review.append("PR: " + describe(res))
    status = "write" if (writes or tag_changes) else ("review" if review else "no_vendor_hit")
    if (writes or tag_changes) and review:
        status = "write+review"
    return {"uuid": u, "street": r.get("street"), "zip": r.get("zip"), "city": r.get("city"),
            "owner": clean_name(f"{r.get('owner_first') or ''} {r.get('owner_last') or ''}"),
            "lists": r.get("lists"), "status": status, "rel_slots": sorted(slots),
            "missing_rel": missing_rel, "pr_tagged": len(pr_phones), "pr_missing_before": pr_missing_before,
            "pr_before": [pr_first, pr_last], "writes": writes, "tag_changes": tag_changes,
            "findings": findings, "review": review}


def cmd_plan(a) -> int:
    records = jload(RECORDS, {})
    idx = jload(INDEX, None)
    if not idx:
        say("no index.json -- run `index` first")
        return 2
    say(f"{len(records)} records, index of {len(idx['persons'])} persons")
    plans, counts = [], {}
    for u, r in records.items():
        if r.get("_gone") or r.get("_error"):
            continue
        p = plan_record(u, r, idx)
        if p is None:
            continue
        counts[p["status"]] = counts.get(p["status"], 0) + 1
        if p["status"] != "complete":
            plans.append(p)
    n_rel = sum(1 for p in plans if p.get("missing_rel"))
    n_pr = sum(1 for p in plans if p.get("pr_missing_before"))
    n_tag = sum(len(p.get("tag_changes") or []) for p in plans)
    say(f"statuses: {counts}")
    say(f"records with unnamed REL slots: {n_rel}; with unnamed/junk PR: {n_pr}; tag corrections: {n_tag}")
    jsave(PLAN, {"built": datetime.now().isoformat(timespec="seconds"), "counts": counts, "plans": plans})
    with open(RUN / "plan.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["status", "street", "zip", "owner", "lists", "missing_rel", "pr_tagged", "pr_before",
                    "writes", "tag_changes", "review", "uuid"])
        for p in sorted(plans, key=lambda p: (p["status"], p["street"] or "")):
            w.writerow([p["status"], p["street"], p["zip"], p["owner"], ";".join(p.get("lists") or []),
                        ";".join(f"REL{n}" for n in p["missing_rel"]), p["pr_tagged"],
                        " | ".join(p["pr_before"]), json.dumps(p["writes"], ensure_ascii=False),
                        ";".join(f"{t['number']}:{t['from']}->{t['to']}" for t in p["tag_changes"]),
                        " || ".join(p["review"]), p["uuid"]])
    say(f"wrote {PLAN} and plan.csv ({len(plans)} rows)")
    return 0


# ----------------------------------------------------------------------------- write

def active_fields(api: Api) -> dict[str, str]:
    """label -> uuid for ACTIVE custom fields only. The account carries an inactive twin of
    'REL1: Full Name' (8752277a...); writing to it would vanish into a dead field."""
    st, b = api.get("/api/internal/custom-fields/?limit=999")
    rows = b if isinstance(b, list) else b.get("results", []) if isinstance(b, dict) else []
    out = {r["label"]: r["uuid"] for r in rows if r.get("is_active") and r.get("label")}
    for must in ("REL1: Full Name", "PR First Name", "PR Last Name"):
        if must not in out:
            raise RuntimeError(f"custom-fields read ({st}) lacks {must!r} -- refusing to write anything")
    return out


def live_state(api: Api, uuid: str) -> dict:
    st, body = api.get(f"/api/internal/property/{uuid}/")
    if st != 200:
        return {"_error": st}
    stc, cf = api.get(f"/api/internal/property/{uuid}/custom-field/")
    if stc != 200:
        return {"_error": f"cf {stc}"}
    owner = body.get("owner") or {}
    return {"cf": cf_values(cf),
            "owner_uuid": owner.get("uuid"),
            "phones": [{"number": p.get("number"), "type": p.get("type"), "tags": tag_names(p.get("tags")),
                        "status": p.get("status"), "is_connected": p.get("is_connected", p.get("isConnected")),
                        "verified": p.get("verified")}
                       for p in owner.get("phones") or []]}


def _phone_payload(ph: dict, tags: list[str]) -> dict:
    """Upsert row for a phone that already exists: everything passed through UNCHANGED
    except the tags. `is_connected` / `verified` are sent only when the record carries
    them (a null must not become True/False on every phone of every owner touched)."""
    out = {"number": ph["number"], "type": ph["type"] or "UNKNOWN", "tags": tags,
           "status": ph["status"] or "UNKNOWN"}
    if ph.get("is_connected") is not None:
        out["is_connected"] = bool(ph["is_connected"])
    if ph.get("verified") is not None:
        out["verified"] = bool(ph["verified"])
    return out


def verify_against_plan(live: dict, p: dict, labels: list[str] | None = None) -> list[str]:
    problems = []
    for lbl, val in p["writes"].items():
        if labels is not None and lbl not in labels:
            continue
        got = live["cf"].get(lbl, "")
        if lbl.endswith("Full Name") or lbl.startswith("PR "):
            if clean_name(got).lower() != clean_name(val).lower():
                problems.append(f"{lbl}: want {val!r} got {got!r}")
        elif d10(got) != d10(val):
            problems.append(f"{lbl}: want {val!r} got {got!r}")
    by_num = {d10(ph["number"]): ph for ph in live["phones"]}
    for t in p["tag_changes"]:
        cur = by_num.get(d10(t["number"]))
        if not cur:
            problems.append(f"phone {t['number']} missing")
        elif t["to"] not in cur["tags"] or t["from"] in cur["tags"]:
            problems.append(f"phone {t['number']} tags {cur['tags']} (want {t['to']} not {t['from']})")
    return problems


def phone_order_changed(before: list[dict], after: list[dict]) -> str:
    """The lane scoring wrote the dial ORDER through the upsert list; a retag must not
    reshuffle it. '' when the sequence of numbers is unchanged."""
    b = [d10(x["number"]) for x in before]
    a = [d10(x["number"]) for x in after]
    if b == a:
        return ""
    return f"phone order changed: {len(b)} -> {len(a)} numbers, first diff at position " \
           f"{next((i for i, (x, y) in enumerate(zip(b, a)) if x != y), min(len(a), len(b)))}"


def write_one(api: Api, p: dict, fields: dict[str, str], commit: bool) -> dict:
    log = {"uuid": p["uuid"], "street": p["street"], "at": datetime.now().isoformat(timespec="seconds"),
           "commit": commit, "steps": {}}
    live = live_state(api, p["uuid"])
    if live.get("_error"):
        log["error"] = f"detail {live['_error']}"
        return log
    # Guard: only write a field that is STILL empty (or still the junk we planned to replace).
    body, skipped, backup_cf = [], [], {}
    pr_junk_now = is_junk_pr(live["cf"].get("PR First Name"), live["cf"].get("PR Last Name"))
    for lbl, val in p["writes"].items():
        cur = clean_name(live["cf"].get(lbl))
        if cur and not (lbl.startswith("PR ") and pr_junk_now):
            skipped.append(f"{lbl} now holds {cur!r}")
            continue
        if lbl not in fields:
            skipped.append(f"{lbl} not an active field")
            continue
        backup_cf[lbl] = live["cf"].get(lbl, "")
        body.append({"field_uuid": fields[lbl], "value": val})
    change_nums = {d10(t["number"]) for t in p["tag_changes"]}
    log["backup"] = {"cf": backup_cf, "owner_uuid": live["owner_uuid"],
                     # the full list in its live order, so undo can restore tags AND order
                     "phones": list(live["phones"]) if change_nums else []}
    log["skipped"] = skipped
    if not commit:
        log["would_write"] = {"cf": [(lbl, p["writes"][lbl]) for lbl in backup_cf],
                              "tag_changes": [(t["number"], t["from"], t["to"]) for t in p["tag_changes"]]}
        return log
    if body:
        st, resp = api.write("PATCH", f"/api/internal/property/{p['uuid']}/custom-field/update-values/", body)
        log["steps"]["cf"] = {"status": st, "sent": len(body), "resp": str(resp)[:200]}
        if st not in (200, 201, 204):
            log["error"] = f"custom-field update {st}"
            return log
    if p["tag_changes"] and live["owner_uuid"]:
        # Send the FULL phone list in its current order with only the planned tags changed:
        # the upsert list order is the dial order the lane scoring wrote (dpd_lane_phone_score),
        # so a one-number upsert could reposition it.
        retag = {}
        live_pr = clean_name(f"{live['cf'].get('PR First Name', '')} {live['cf'].get('PR Last Name', '')}")
        used_pr_n = {(parse_tag(t) or ("", None, None))[1] for ph in live["phones"] for t in ph["tags"]
                     if (parse_tag(t) or ("",))[0] == "pr"}
        for t in p["tag_changes"]:
            cur = next((ph for ph in live["phones"] if d10(ph["number"]) == d10(t["number"])), None)
            if not cur or t["from"] not in cur["tags"]:
                log["steps"].setdefault("tag_skipped", []).append(f"{t['number']} no longer carries {t['from']}")
                continue
            if live_pr and not is_junk_pr(live["cf"].get("PR First Name"), live["cf"].get("PR Last Name")) \
                    and not names_match(t.get("who") or "", live_pr):
                log["steps"].setdefault("tag_skipped", []).append(
                    f"{t['number']}: live PR is {live_pr!r}, vendor says {t.get('who')!r} -- not retagged")
                continue
            to = t["to"]
            if to in cur["tags"]:
                # the number already carries the Pr tag; only the Rel tag comes off
                retag[d10(t["number"])] = [x for x in cur["tags"] if x != t["from"]]
                continue
            k = t.get("to_n") or 1
            while k in used_pr_n:   # a Pr.N added since the plan
                k += 1
                to = f"Pr.{k}"
            used_pr_n.add(k)
            t["to"] = to
            retag[d10(t["number"])] = [x for x in cur["tags"] if x != t["from"]] + [to]
        if retag:
            payload = [_phone_payload(ph, retag.get(d10(ph["number"]), ph["tags"])) for ph in live["phones"]]
            st, resp = api.write("POST", f"/api/internal/owner/{live['owner_uuid']}/upsert-phones/",
                                 {"phones": payload})
            log["steps"]["phones"] = {"status": st, "sent": len(payload), "retagged": len(retag),
                                      "resp": str(resp)[:200],
                                      "changes": [(n, tags) for n, tags in retag.items()]}
            if st not in (200, 201, 204):
                log["error"] = f"upsert-phones {st}"
                return log
    after = live_state(api, p["uuid"])
    if after.get("_error"):
        log["error"] = f"read-back failed ({after['_error']}) -- write may have landed; check by hand"
        return log
    log["verify"] = verify_against_plan(after, p, labels=list(backup_cf))
    if p["tag_changes"] and log["steps"].get("phones"):
        order = phone_order_changed(live["phones"], after["phones"])
        if order:
            log["verify"].append(order)
        # full tag-set equality: retagged numbers = before - from + to; every other number
        # keeps exactly the tags it had
        after_by = {d10(ph["number"]): set(ph["tags"]) for ph in after["phones"]}
        for ph in live["phones"]:
            want = set(retag.get(d10(ph["number"]), ph["tags"]))
            got = after_by.get(d10(ph["number"]))
            if got is None:
                log["verify"].append(f"phone {ph['number']} vanished after upsert")
            elif got != want:
                log["verify"].append(f"phone {ph['number']} tags {sorted(got)} != expected {sorted(want)}")
        flags_before = {d10(ph["number"]): (ph.get("is_connected"), ph.get("verified")) for ph in live["phones"]}
        for ph in after["phones"]:
            if flags_before.get(d10(ph["number"])) not in (None, (ph.get("is_connected"), ph.get("verified"))):
                log["verify"].append(f"phone {ph['number']} flags changed {flags_before[d10(ph['number'])]} -> "
                                     f"{(ph.get('is_connected'), ph.get('verified'))}")
    inactive = [s for s in skipped if "not an active field" in s]
    if inactive:
        log["verify"].append(f"{len(inactive)} planned field(s) not active: {inactive}")
    log["ok"] = not log["verify"]
    log["wrote_anything"] = bool(body) or bool(log["steps"].get("phones"))
    return log


def _load_plans(a) -> list[dict]:
    plan = jload(PLAN, None)
    if not plan:
        raise SystemExit("no plan.json -- run `plan` first")
    plans = [p for p in plan["plans"] if p["status"] in ("write", "write+review")]
    if getattr(a, "key", None):
        plans = [p for p in plans if a.key.lower() in (p["street"] or "").lower() or a.key == p["uuid"]]
    if getattr(a, "limit", 0):
        plans = plans[:a.limit]
    return plans


def cmd_probe(a) -> int:
    plans = _load_plans(a)
    if len(plans) != 1:
        say(f"--key must select exactly one record (got {len(plans)})")
        return 2
    api = Api()
    fields = active_fields(api)
    p = plans[0]
    say(json.dumps({k: p[k] for k in ("street", "owner", "writes", "tag_changes", "review")},
                   indent=1, ensure_ascii=False))
    log = write_one(api, p, fields, commit=False)
    say("DRY: " + json.dumps(log.get("would_write"), ensure_ascii=False) + f" skipped={log['skipped']}")
    if not a.commit:
        return 0
    log = write_one(api, p, fields, commit=True)
    ts = datetime.now().strftime("%Y%m%dT%H%M%S")
    jsave(RUN / f"backup_probe_{ts}.json", [log])
    say(json.dumps(log, indent=1, ensure_ascii=False, default=str))
    return 0 if log.get("ok") else 1


def cmd_commit(a) -> int:
    plans = _load_plans(a)
    if a.skip_done:
        done_ok = {json.loads(l)["uuid"] for f in RUN.glob("writelog_*.jsonl")
                   for l in f.read_text(encoding="utf-8").splitlines()
                   if l.strip() and json.loads(l).get("ok") and json.loads(l).get("commit")}
        plans = [p for p in plans if p["uuid"] not in done_ok]
        say(f"skipping {len(done_ok)} records already written ok")
    say(f"{len(plans)} records to write ({'COMMIT' if a.commit else 'dry run'})")
    if not plans:
        say("FAIL: nothing selected -- an empty plan is not a clean run")
        return 2
    api = Api()
    fields = active_fields(api)
    ts = datetime.now().strftime("%Y%m%dT%H%M%S")
    backup = RUN / f"backup_{ts}.json"
    log_path = RUN / f"writelog_{ts}.jsonl"
    logs, ok, bad = [], 0, 0
    with open(log_path, "a", encoding="utf-8") as lf:
        for i, p in enumerate(plans, 1):
            log = write_one(api, p, fields, commit=a.commit)
            logs.append(log)
            lf.write(json.dumps(log, ensure_ascii=False, default=str) + "\n")
            lf.flush()
            if log.get("ok"):
                ok += 1
            elif a.commit:
                bad += 1
                say(f"  PROBLEM {p['street']}: {log.get('error') or log.get('verify')}")
            if i % 25 == 0 or i == len(plans):
                jsave(backup, logs)
                say(f"  {i}/{len(plans)} ok={ok} bad={bad}")
    jsave(backup, logs)
    say(f"{'committed' if a.commit else 'dry run'}: ok={ok} bad={bad}; backup {backup}")
    return 0 if bad == 0 else 1


def skipped_labels_from_logs() -> dict[str, set[str]]:
    """{uuid: labels the writer skipped because another source had filled them} -- QA must
    not report those as failures."""
    out: dict[str, set[str]] = {}
    for f in RUN.glob("writelog_*.jsonl"):
        for l in f.read_text(encoding="utf-8").splitlines():
            if not l.strip():
                continue
            log = json.loads(l)
            for s in log.get("skipped") or []:
                if " now holds " in s:
                    out.setdefault(log["uuid"], set()).add(s.split(" now holds ")[0])
    return out


def cmd_qa(a) -> int:
    plans = _load_plans(a)
    skipped = skipped_labels_from_logs()
    api = Api()
    fails = []
    for i, p in enumerate(plans, 1):
        live = live_state(api, p["uuid"])
        if live.get("_error"):
            fails.append((p["street"], f"detail {live['_error']}"))
            continue
        labels = [lbl for lbl in p["writes"] if lbl not in skipped.get(p["uuid"], set())]
        probs = verify_against_plan(live, p, labels=labels)
        if probs:
            fails.append((p["street"], probs))
        if i % 50 == 0:
            say(f"  {i}/{len(plans)} fails={len(fails)}")
    say(f"QA: {len(plans) - len(fails)}/{len(plans)} clean")
    for s, pr in fails[:50]:
        say(f"  FAIL {s}: {pr}")
    jsave(RUN / f"qa_{datetime.now().strftime('%Y%m%dT%H%M%S')}.json",
          {"clean": len(plans) - len(fails), "fails": fails})
    return 0 if not fails else 1


def cmd_undo(a) -> int:
    """Restore prior field values / phone tags from a backup file. Only what THIS run
    actually wrote (a 2xx step) and that still holds our value is touched -- a field another
    plan filled since is left alone. Exit 1 on any non-2xx."""
    logs = jload(Path(a.backup), [])
    plan = {p["uuid"]: p for p in (jload(PLAN, {}).get("plans") or [])}
    api = Api()
    fields = active_fields(api)
    n = bad = 0
    for log in logs:
        if not log.get("commit") or not log.get("backup"):
            continue
        steps = log.get("steps") or {}
        b = log["backup"]
        live = live_state(api, log["uuid"])
        if live.get("_error"):
            say(f"  {log['street']}: cannot read record ({live['_error']}); skipped")
            bad += 1
            continue
        written = (plan.get(log["uuid"]) or {}).get("writes") or {}
        body = []
        if (steps.get("cf") or {}).get("status") in (200, 201, 204):
            for lbl, prior in b["cf"].items():
                if lbl not in fields:
                    continue
                now = clean_name(live["cf"].get(lbl))
                if now.lower() != clean_name(written.get(lbl, "")).lower():
                    say(f"  {log['street']}: {lbl} now {now!r}, not ours -- left alone")
                    continue
                body.append({"field_uuid": fields[lbl], "value": prior})
        if body:
            st, resp = api.write("PATCH", f"/api/internal/property/{log['uuid']}/custom-field/update-values/", body)
            say(f"  {log['street']}: cf restore {st} ({len(body)} fields)")
            if st not in (200, 201, 204):
                bad += 1
            else:
                back = live_state(api, log["uuid"])
                wrong = [lbl for lbl, prior in b["cf"].items()
                         if lbl in fields and clean_name((back.get("cf") or {}).get(lbl)) != clean_name(prior)
                         and any(x["field_uuid"] == fields[lbl] for x in body)]
                if wrong:
                    say(f"  {log['street']}: read-back after restore differs on {wrong}")
                    bad += 1
        if b.get("phones") and b.get("owner_uuid") and (steps.get("phones") or {}).get("status") in (200, 201, 204):
            payload = [_phone_payload(ph, ph["tags"]) for ph in b["phones"]]
            st, resp = api.write("POST", f"/api/internal/owner/{b['owner_uuid']}/upsert-phones/", {"phones": payload})
            say(f"  {log['street']}: phone tags restore {st}")
            if st not in (200, 201, 204):
                bad += 1
        n += 1
    say(f"undo: {n} records processed, {bad} problems")
    return 0 if bad == 0 else 1


# ----------------------------------------------------------------------------- dial sheets

def cmd_regen_sheets(a) -> int:
    """Re-post the `WHO YOU ARE DIALING` dial sheet on every repaired record that already
    carries one (the Re-Skiptrace review's 17), so each number now shows its name.

    The sheet text in reskiptrace_review's plan.json is frozen at plan time, so this
    rebuilds it from LIVE data (record_context reads the REL/PR fields this repair just
    wrote) with the review's own decision, and posts through its poster, which replaces
    only the same family and logs what it deleted.
    """
    import reskiptrace_review as RR
    review_plan = jload(RR.PLAN, {})
    push_log = jload(RR.PUSH_LOG, {})
    plans = {p["uuid"]: p for p in _load_plans(a)}
    if a.all_sheets:
        targets = [u for u in push_log if u in review_plan and not push_log[u].get("error")]
    else:
        targets = [u for u in push_log if u in plans and u in review_plan and not push_log[u].get("error")]
    say(f"{len(targets)} records carry a dial sheet and were repaired")
    api = Api()
    out = []
    for u in targets:
        dec = review_plan[u]
        st, d = api.get(f"/api/internal/property/{u}/")
        if st != 200:
            say(f"  {dec['street']}: detail {st}")
            continue
        ctx = RR.record_context(api, u, d)
        note = RR.build_note(dec, ctx, dec.get("principals"), dec.get("owner_set"))
        unnamed = note.count("name NOT on file")
        if not a.commit:
            say(f"  DRY {dec['street']}: {len(ctx['phones'])} numbers, {unnamed} still unnamed")
            out.append({"uuid": u, "street": dec["street"], "unnamed": unnamed, "note": note})
            continue
        res = RR.post_pinned_note(api, u, note)
        ok = not res.get("error") and res.get("pin") in (200, 201, 204)
        say(f"  {'ok ' if ok else 'ERR'} {dec['street']}: post={res.get('status')} pin={res.get('pin')} "
            f"replaced={len(res.get('replaced') or [])} unnamed_left={unnamed}")
        out.append({"uuid": u, "street": dec["street"], "unnamed": unnamed, "result": res, "note": note})
    jsave(RUN / f"regen_sheets_{datetime.now().strftime('%Y%m%dT%H%M%S')}.json", out)
    return 0


# ----------------------------------------------------------------------------- main

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("pull")
    p.add_argument("--fresh", action="store_true")
    p.add_argument("--threads", type=int, default=4)
    p.add_argument("--limit", type=int, default=0)
    p.set_defaults(fn=cmd_pull)
    p = sub.add_parser("index")
    p.set_defaults(fn=cmd_index)
    p = sub.add_parser("plan")
    p.set_defaults(fn=cmd_plan)
    p = sub.add_parser("probe")
    p.add_argument("--key", required=True)
    p.add_argument("--commit", action="store_true")
    p.set_defaults(fn=cmd_probe)
    p = sub.add_parser("commit")
    p.add_argument("--commit", action="store_true", help="write; without it a dry run")
    p.add_argument("--key", default="")
    p.add_argument("--limit", type=int, default=0)
    p.add_argument("--skip-done", action="store_true", help="skip records an earlier writelog marked ok")
    p.set_defaults(fn=cmd_commit)
    p = sub.add_parser("qa")
    p.add_argument("--key", default="")
    p.add_argument("--limit", type=int, default=0)
    p.set_defaults(fn=cmd_qa)
    p = sub.add_parser("undo")
    p.add_argument("--backup", required=True)
    p.set_defaults(fn=cmd_undo)
    p = sub.add_parser("regen-sheets")
    p.add_argument("--commit", action="store_true")
    p.add_argument("--all-sheets", action="store_true", help="every pushed sheet, not only repaired records")
    p.set_defaults(fn=cmd_regen_sheets)
    a = ap.parse_args()
    return a.fn(a)


if __name__ == "__main__":
    sys.exit(main())
